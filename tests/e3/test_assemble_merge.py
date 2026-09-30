"""Character patches, overlap preservation, and the saved 24-edit all-adopt regression."""
from __future__ import annotations

import copy
import io
import json
from pathlib import Path

import pytest

from neumann.analyze import assemble as asm


def edits(before, *after):
    return [{"edit_id": f"e{i}", "card_id": f"c{i}", "plan_line": 1, "kind": "replace",
             "current_text": before, "proposed_text": text,
             "rationale": {"text": f"rationale {i}", "excerpt_ids": [f"ex{i}"]}}
            for i, text in enumerate(after)]


def adopt(items):
    return [{"edit_id": item["edit_id"], "decision": "adopt"} for item in items]


@pytest.fixture
def live24():
    return json.loads(Path(__file__).with_name("assemble_live_smoke24.json").read_text(encoding="utf-8"))


def test_nonoverlapping_patches_recalculate_offsets_and_keep_all_sources():
    before = "alpha / beta / gamma"
    items = edits(before, "longer alpha / beta / gamma", "alpha / short / gamma", "alpha / beta / final gamma")
    decisions = adopt(items)
    out = asm.assemble_revised_plan(before, items, decisions)
    assert out["revised_text"] == "longer alpha / short / final gamma"
    assert out["stats"]["merged"] == out["stats"]["applied"] == 3
    assert out["lines"][0]["edit_ids"] == ["e0", "e1", "e2"]
    assert {c["edit_id"]: c["excerpt_ids"] for c in out["changes"]} == {"e0": ["ex0"], "e1": ["ex1"], "e2": ["ex2"]}
    ev = {f"ex{i}": {"text": f"source {i}", "source_url": "https://example.org"} for i in range(3)}
    md = asm.render_markdown(out, ev)
    assert md["footnoted"].startswith(out["revised_text"] + "[^1][^2][^3]")
    from docx import Document
    doc = Document(io.BytesIO(asm.build_docx(out, ev)))
    assert [r.text for p in doc.paragraphs for r in p.runs if r.font.superscript] == ["1", "2", "3"]
    decisions.reverse()
    again = asm.assemble_revised_plan(before, items, decisions)
    assert {k: v for k, v in out.items() if k != "generated_at"} == {k: v for k, v in again.items() if k != "generated_at"}
    assert asm.validate_revised_plan(out) == []


@pytest.mark.parametrize("after", [("first sentence", "second sentence"), ("alpha new", "alpha alternate"), ("alpha first beta", "alpha second beta")])
def test_overlap_and_same_position_insertions_are_preserved(after):
    before = "alpha beta"
    items = edits(before, *after)
    out = asm.assemble_revised_plan(before, items, adopt(items))
    assert out["revised_text"].splitlines() == list(after)
    assert out["stats"]["applied"] == 2 and out["stats"]["converted_insert"] == 1
    assert out["changes"][1]["effective_kind"] == "insert_after"
    assert out["changes"][1]["kind"] == "replace"
    assert out["edit_statuses"][1]["reason"]


def test_every_unapplied_adoption_has_a_reason():
    items = edits("original", "one", "two", "", "four")
    items[0]["current_text"] = "stale"
    items[1]["plan_line"] = 99
    items[3]["kind"] = "delete"
    out = asm.assemble_revised_plan("original", items, adopt(items) + [{"edit_id": "unknown", "decision": "adopt"}])
    assert out["stats"]["applied"] == 0 and out["stats"]["not_applied"] == 5
    assert len(out["edit_statuses"]) == 5
    assert all(s["status"] == "not_applied" and s["reason"] for s in out["edit_statuses"])
    assert asm.validate_revised_plan(out) == []
    assert "편집별 적용 상태" in asm.render_markdown(out, {})["history"]
    from docx import Document
    doc = Document(io.BytesIO(asm.build_docx(out, {})))
    cells = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert all(s["edit_id"] in cells and s["reason"] in cells for s in out["edit_statuses"])


def test_live_24_all_adopt_zero_silent_omissions(live24):
    before = copy.deepcopy(live24)
    revision = live24["revision"]
    items = [e for r in revision["revisions"] for e in r["edits"]]
    out = asm.assemble_revised_plan(live24["plan_text"], revision, adopt(items), result=live24["result"], regate=True)
    assert out["stats"]["applied"] == 24
    assert out["stats"]["conflicts"] == out["stats"]["not_applied"] == out["stats"]["skipped"] == 0
    assert len(out["edit_statuses"]) == len(out["changes"]) == 24
    by_edit = {c["edit_id"]: c for c in out["changes"]}
    for item in items:
        change = by_edit[item["edit_id"]]
        assert change["applied_text"] == item["proposed_text"]
        assert change["new_text"] in out["revised_text"]
        assert change["excerpt_ids"] == item["rationale"]["excerpt_ids"]
        assert change["rationale"] == item["rationale"]["text"]
    assert out["stats"]["converted_insert"] == 11
    assert live24 == before
    assert asm.validate_revised_plan(out) == []


def test_polish_of_combined_line_cannot_drop_a_contribution():
    items = edits("alpha / beta", "alpha revised / beta", "alpha / beta reviewed")
    out = asm.assemble_revised_plan("alpha / beta", items, adopt(items))
    assert out["stats"]["merged"] == 2
    def drop(schema, instructions, payload, **kwargs):
        return {"lines": [{"no": 1, "text": "alpha revised / beta"}]}
    polished = asm.polish_revised_plan(out, drop)
    assert polished["polish"]["applied"] is False
    assert polished["revised_text"] == out["revised_text"]
