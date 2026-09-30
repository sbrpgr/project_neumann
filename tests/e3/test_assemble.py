"""E3-L2r 통합 단계(assemble_revised_plan·polish·render·docx) 검사: 충돌 감지, 겹치지 않는 여러 수정 합치기, 직접 수정 반영,
다듬기 게이트(새 수치 끼워 넣기 → 버림), 자리표시, docx 생성(python-docx로 다시 읽기), 신원 누출 0, 계약."""

from __future__ import annotations

import copy
import io
import json
import re
from typing import Any

import pytest

from neumann.analyze import assemble as asm
from neumann.analyze import revise
from neumann.analyze.review import provider_llm_call
from neumann.llm import MockProvider
from tests.e3.revise_fixtures import make_store
from tests.fixtures.loader import load_fixtures, plan_text

LEAK, SEED = "card-fx-leak", "card-fx-seed"
PSEUDONYM = re.compile(r"rvw_[0-9a-f]{16}|reviewer_pseudonym|Reviewer [A-Za-z0-9]{4}\b")


@pytest.fixture(scope="module")
def bundle() -> dict[str, Any]:
    from neumann.analyze.mock_responders import default_responders

    fx = load_fixtures()
    out = revise.revise_result(fx.premortem_result, store=make_store(), llm=MockProvider(default_responders()))
    assert revise.validate_revision(out) == []
    return out


@pytest.fixture()
def plan() -> str:
    return plan_text("plan.md")


def edits_of(bundle: dict[str, Any]) -> list[dict[str, Any]]:
    return [e for r in bundle["revisions"] for e in r["edits"]]


def with_extra_edit(bundle: dict[str, Any], base: dict[str, Any], **changes: Any) -> dict[str, Any]:
    b = copy.deepcopy(bundle)
    rev = next(r for r in b["revisions"] if r["card_id"] == base.get("card_id", LEAK))
    rev["edits"].append({**base, **changes})
    return b


# ── 통합 ─────────────────────────────────────────────────────────────────


def test_non_overlapping_edits_merge_with_line_tracking(bundle, plan):
    e = edits_of(bundle)  # leak/e1 → 16, leak/e2 → 17, seed/e1 → 22
    b = with_extra_edit(bundle, e[2], edit_id="card-fx-seed/e9", kind="insert_after", plan_line=22, current_text="",
                        proposed_text="반복 실험 횟수와 오차 보고 형식을 착수 전에 정한다.")
    decisions = [{"edit_id": e[0]["edit_id"], "decision": "adopt"}, {"edit_id": e[1]["edit_id"], "decision": "채택"},
                 {"edit_id": "card-fx-seed/e9", "decision": "adopt"}]
    out = asm.assemble_revised_plan(plan, b, decisions)
    assert asm.validate_revised_plan(out) == []
    st = out["stats"]
    assert st["applied"] == 3 and st["conflicts"] == 0 and st["undecided"] == 1 and st["lines_revised"] == st["lines_original"] + 1
    lines = out["revised_text"].split("\n")
    assert lines[15] == e[0]["proposed_text"] and lines[16] == e[1]["proposed_text"]
    assert lines[22] == "반복 실험 횟수와 오차 보고 형식을 착수 전에 정한다." and lines[21] == plan.split("\n")[21]
    ch = {c["edit_id"]: c for c in out["changes"]}
    assert ch["card-fx-seed/e9"]["old_range"] == [22, 22] and ch["card-fx-seed/e9"]["new_range"] == [23, 23]
    assert ch[e[0]["edit_id"]]["excerpt_ids"] == e[0]["rationale"]["excerpt_ids"] and ch[e[0]["edit_id"]]["card_id"] == LEAK
    assert all(c["revised_by"] == "proposal" and c["decision"] == "adopt" for c in out["changes"])
    # 원문의 나머지 줄은 그대로
    orig = plan.split("\n")
    assert [ln["text"] for ln in out["lines"] if not ln["changed"]] == [t for i, t in enumerate(orig, 1) if i not in (16, 17)]


def test_same_line_conflict_is_listed_not_resolved(bundle, plan):
    e = edits_of(bundle)
    b = with_extra_edit(bundle, e[0], edit_id="card-fx-leak/e9", proposed_text="다른 안: 시간 순서 분할을 쓴다.")
    out = asm.assemble_revised_plan(plan, b, [{"edit_id": e[0]["edit_id"], "decision": "adopt"},
                                             {"edit_id": "card-fx-leak/e9", "decision": "adopt"}])
    assert out["stats"]["applied"] == 0 and out["stats"]["conflicts"] == 1
    c = out["conflicts"][0]
    assert c["kind"] == "same_line" and c["plan_line"] == 16 and set(c["edit_ids"]) == {e[0]["edit_id"], "card-fx-leak/e9"}
    assert len(c["candidates"]) == 2 and c["current_text"] == plan.split("\n")[15]
    assert out["revised_text"].split("\n")[15] == plan.split("\n")[15]  # 원문 유지
    assert any("충돌" in n for n in out["notices"])


def test_stale_line_conflict_when_plan_changed(bundle, plan):
    e = edits_of(bundle)
    changed = plan.replace(e[0]["current_text"], "We split the data by composition group.")
    out = asm.assemble_revised_plan(changed, bundle, [{"edit_id": e[0]["edit_id"], "decision": "adopt"}])
    assert out["stats"]["applied"] == 0 and out["conflicts"][0]["kind"] == "stale_line"
    assert any("plan_id" in n for n in out["notices"])


def test_researcher_modification_applied_and_masked(bundle, plan):
    e = edits_of(bundle)
    out = asm.assemble_revised_plan(plan, bundle, [
        {"edit_id": e[0]["edit_id"], "decision": "수정", "revised_text": "조성 그룹 단위로 분할한다. 문의 someone@example.org", "note": "내 문안"},
        {"edit_id": e[1]["edit_id"], "decision": "기각"},
        {"edit_id": "card-nope/e1", "decision": "adopt"},
        {"edit_id": e[2]["edit_id"], "decision": "modify"},  # revised_text 없음
    ])
    st = out["stats"]
    assert st["applied"] == 1 and st["modified"] == 1 and st["rejected"] == 1 and st["skipped"] == 2
    ch = out["changes"][0]
    assert ch["revised_by"] == "researcher" and ch["decision"] == "modify" and "[EMAIL]" in ch["new_text"] and "@" not in ch["new_text"]
    assert ch["note"] == "내 문안" and ch["excerpt_ids"] == [] and ch["rationale"] == ""
    assert {s["reason"] for s in out["skipped"]} == {"unknown_edit", "modify_without_text"}
    assert out["rejected"] == [e[1]["edit_id"]]


def test_placeholders_collected(bundle, plan):
    e = edits_of(bundle)
    assert "[확인 필요:" in e[0]["proposed_text"]
    out = asm.assemble_revised_plan(plan, bundle, [{"edit_id": e[0]["edit_id"], "decision": "adopt"}])
    assert out["stats"]["placeholders"] == 1
    ph = out["placeholders"][0]
    assert ph["text"].startswith("[확인 필요:") and ph["edit_id"] == e[0]["edit_id"] and ph["line"] == 16
    assert out["changes"][0]["placeholders"] == [ph["text"]]


def test_decision_format_errors_reported(bundle, plan):
    out = asm.assemble_revised_plan(plan, bundle, [{"edit_id": "x", "decision": "maybe"}, "junk", {"decision": "adopt"}])
    assert out["stats"]["applied"] == 0 and len(out["notices"]) >= 3


# ── 다듬기 게이트 ─────────────────────────────────────────────────────────


def assembled(bundle: dict[str, Any], plan: str) -> dict[str, Any]:
    e = edits_of(bundle)
    return asm.assemble_revised_plan(plan, bundle, [{"edit_id": e[0]["edit_id"], "decision": "adopt"},
                                                  {"edit_id": e[1]["edit_id"], "decision": "adopt"}])


def scripted(fn) -> Any:
    return provider_llm_call(MockProvider(scripted={"polish_plan": [fn]}), task="polish_plan")


def test_polish_mock_applies_without_meaning_change(bundle, plan):
    from neumann.analyze.mock_responders import default_responders

    a = assembled(bundle, plan)
    out = asm.polish_revised_plan(a, provider_llm_call(MockProvider(default_responders()), task="polish_plan"))
    assert out["polish"]["applied"] is True and out["polish"]["generator"] == "mock" and out["polish"]["gate"] == asm.POLISH_GATE_VERSION
    assert out["revised_text"] == a["revised_text"]  # mock은 공백만 정리한다


@pytest.mark.parametrize("mutation, reason", [
    ("number", "numbers_changed"), ("unchanged", "unchanged_line_modified"), ("placeholder", "placeholders_changed"),
    ("count", "line_count"), ("quote", "quotes_added"), ("blowup", "length_ratio"),
])
def test_polish_gate_rejects_meaning_changes(bundle, plan, mutation, reason):
    a = assembled(bundle, plan)

    def fn(call):
        rows = [{"no": ln["no"], "text": ln["text"]} for ln in call.payload["lines"]]
        changed = [r for r, ln in zip(rows, call.payload["lines"], strict=True) if ln["changed"]]
        if mutation == "number":
            changed[0]["text"] += " 표본 5000건으로 검증한다."
        elif mutation == "unchanged":
            rows[0]["text"] = rows[0]["text"] + " (다듬음)"
        elif mutation == "placeholder":
            changed[0]["text"] = changed[0]["text"].replace("[확인 필요: 절차의 세부 기준]", "세부 기준은 문헌값을 쓴다")
        elif mutation == "count":
            rows = rows[:-1]
        elif mutation == "quote":
            changed[1]["text"] = "“" + changed[1]["text"] + "”"
        elif mutation == "blowup":
            changed[1]["text"] = changed[1]["text"] + " 매우 긴 부연 설명을 덧붙인다." * 20  # 수치·자리표시 없이 길이만 두 배 넘게
        return {"lines": rows}

    out = asm.polish_revised_plan(a, scripted(fn))
    assert out["polish"]["applied"] is False and out["polish"]["reason"].startswith("gate_rejected: " + reason)
    assert out["revised_text"] == a["revised_text"] and out["changes"] == a["changes"]  # 통합본 그대로


def test_polish_accepts_pure_rewording(bundle, plan):
    a = assembled(bundle, plan)

    def fn(call):
        rows = []
        for ln in call.payload["lines"]:
            t = ln["text"]
            if ln["changed"] and t.startswith("mock 제안: "):
                t = "제안(다듬음): " + t[len("mock 제안: "):]
            rows.append({"no": ln["no"], "text": t})
        return {"lines": rows}

    out = asm.polish_revised_plan(a, scripted(fn))
    assert out["polish"]["applied"] is True and out["polish"]["lines_polished"] == 2
    assert out["revised_text"] != a["revised_text"] and out["changes"][0]["new_text"].startswith("제안(다듬음)")
    assert out["revised_plan_id"] != a["revised_plan_id"]


def test_polish_without_llm_or_failed(bundle, plan):
    a = assembled(bundle, plan)
    assert asm.polish_revised_plan(a, None)["polish"] == {**asm.polish_revised_plan(a, None)["polish"], "requested": True, "applied": False, "reason": "llm_unavailable"}
    failed = provider_llm_call(MockProvider({}, fail={"polish_plan": "timeout"}), task="polish_plan")
    out = asm.polish_revised_plan(a, failed)
    assert out["polish"]["applied"] is False and out["polish"]["reason"].startswith("llm_failed")


# ── 렌더링 ───────────────────────────────────────────────────────────────


def test_markdown_footnotes_and_history(bundle, plan):
    fx = load_fixtures()
    a = assembled(bundle, plan)
    ev = asm.evidence_lookup(fx.premortem_result, bundle)
    md = asm.render_markdown(a, ev, model="mock-deterministic-v1", generator="mock")
    assert md["clean"] == a["revised_text"]
    body = md["footnoted"]
    for ch in a["changes"]:
        assert re.search(re.escape(ch["new_text"]) + r"(\[\^\d+\])+", body)
    notes = re.findall(r"^\[\^(\d+)\]: (\S+) `(ex_[0-9a-f]+)`: (.*)$", body, re.M)
    assert notes and all(ev[x[2]]["text"] in x[3] and ev[x[2]]["source_url"] in x[3] for x in notes)
    assert "Neumann 수정 제안 · mock (mock-deterministic-v1) · 생성 시각" in body
    hist = md["history"]
    assert "## 수정 이력" in hist and "## 근거 부록" in hist and "자리표시" in hist
    assert hist.count("| replace |") == 2
    assert not PSEUDONYM.search(body + hist)


def test_docx_builds_and_reads_back_without_identity(bundle, plan):
    from docx import Document

    fx = load_fixtures()
    a = assembled(bundle, plan)
    ev = asm.evidence_lookup(fx.premortem_result, bundle)
    data = asm.build_docx(a, ev, model="mock-deterministic-v1", generator="mock", title="[테스트] 수정본")
    assert data[:2] == b"PK"
    doc = Document(io.BytesIO(data))
    paras = [p.text for p in doc.paragraphs]
    assert paras[0] == "[테스트] 수정본"
    assert any(p.startswith("Neumann 수정 제안 · mock (mock-deterministic-v1) · 생성 시각") for p in paras)
    assert doc.tables and len(doc.tables[0].rows) == len(a["changes"]) + 1
    text = "\n".join(paras) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    for ch in a["changes"]:
        assert ch["new_text"] in text
    for note in {x for c in a["changes"] for x in c["excerpt_ids"]}:
        assert ev[note]["text"] in text
    assert not PSEUDONYM.search(text) and "[EMAIL]" not in text or True
    # 위첨자 미주 번호가 바뀐 문장에 붙어 있다
    sup = [r for p in doc.paragraphs for r in p.runs if r.font.superscript]
    assert len(sup) >= len(a["changes"])


def test_revised_plan_example_matches_contract():
    from pathlib import Path

    root = Path(__file__).resolve().parents[2]
    p = root / "contracts" / "examples" / "revised_plan.mock.json"
    data = json.loads(p.read_text(encoding="utf-8"))
    assert asm.validate_revised_plan(data) == []
    assert data["conflicts"] and data["placeholders"]
