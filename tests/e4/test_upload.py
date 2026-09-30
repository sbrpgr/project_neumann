"""E4-L1a 계획서 업로드 파서 테스트.

pdf·docx 견본은 테스트 안에서 만든다(파일 커밋 없음). PDF는 글자마다 2바이트 CID + ToUnicode CMap을 붙인
최소 PDF를 직접 써서, 한국어가 pypdf 추출을 거쳐 글자 그대로 돌아오는지 본다.
"""

from __future__ import annotations

import builtins
import io
import os
import subprocess
import sys
import tempfile
import threading
import time
import unicodedata
import zipfile
from pathlib import Path

import pytest
import starlette.formparsers
from fastapi import FastAPI, File, UploadFile
from fastapi.testclient import TestClient

import neumann.api.upload as upload
from neumann.api.upload import (
    HWP_MESSAGE,
    MAX_PDF_PAGES,
    MAX_PLAN_CHARS,
    MAX_UPLOAD_BYTES,
    MAX_ZIP_UNCOMPRESSED,
    UploadRejected,
    extract_plan,
    extract_plan_isolated,
    parse_plan_upload,
    router,
)
from neumann.models import PlanDocument

PLAN_MD = (Path(__file__).resolve().parents[1] / "fixtures" / "plans" / "plan.md").read_text(encoding="utf-8")
PLAN_BODY = PLAN_MD.replace("\r\n", "\n").strip("\n")
# CP949에 없는 글자(—)만 바꾼 한국어 본문
PLAN_CP949 = PLAN_BODY.replace("—", "-")
KOREAN_LINES = [
    "연구계획서 — 전해액 이온전도도 예측 대리모델",
    "1. 목표: GNN으로 이온전도도(mS/cm)를 예측한다.",
    "2. 데이터: 문헌 12,000건, 자체 실험 300건",
    "3. 평가: 5겹 교차검증, MAE·R² 보고",
]


# ── 견본 생성 ───────────────────────────────────────────────────────────────


def _stream(payload: bytes) -> bytes:
    return b"<< /Length %d >>\nstream\n" % len(payload) + payload + b"\nendstream"


def make_pdf(pages: list[list[str]]) -> bytes:
    """쪽마다 줄 목록을 받아 최소 PDF를 만든다. 빈 목록이면 텍스트 없는 쪽(스캔 이미지 흉내)."""
    chars = sorted({ch for page in pages for line in page for ch in line})
    blocks = ""
    for i in range(0, len(chars), 100):
        chunk = chars[i : i + 100]
        blocks += f"{len(chunk)} beginbfchar\n" + "".join(f"<{ord(c):04X}> <{ord(c):04X}>\n" for c in chunk) + "endbfchar\n"
    to_unicode = (
        "/CIDInit /ProcSet findresource begin\n12 dict begin\nbegincmap\n"
        "/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n"
        "/CMapName /Adobe-Identity-UCS def\n/CMapType 2 def\n"
        "1 begincodespacerange\n<0000> <FFFF>\nendcodespacerange\n" + blocks + "endcmap\n"
        "CMapName currentdict /CMap defineresource pop\nend\nend\n"
    )
    page_ids = [7 + 2 * i for i in range(len(pages))]
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{p} 0 R' for p in page_ids)}] /Count {len(pages)} >>".encode(),
        b"<< /Type /Font /Subtype /Type0 /BaseFont /TestCJK /Encoding /Identity-H "
        b"/DescendantFonts [4 0 R] /ToUnicode 6 0 R >>",
        b"<< /Type /Font /Subtype /CIDFontType2 /BaseFont /TestCJK "
        b"/CIDSystemInfo << /Registry (Adobe) /Ordering (Identity) /Supplement 0 >> "
        b"/FontDescriptor 5 0 R /DW 1000 /CIDToGIDMap /Identity >>",
        b"<< /Type /FontDescriptor /FontName /TestCJK /Flags 4 /FontBBox [0 -200 1000 900] "
        b"/ItalicAngle 0 /Ascent 900 /Descent -200 /CapHeight 700 /StemV 80 >>",
        _stream(to_unicode.encode("ascii")),
    ]
    for pid, lines in zip(page_ids, pages, strict=True):
        content = ""
        if lines:
            shown = "".join("<" + "".join(f"{ord(c):04X}" for c in line) + "> Tj T*\n" for line in lines)
            content = "BT /F1 10 Tf 14 TL 40 800 Td\n" + shown + "ET\n"
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {pid + 1} 0 R >>".encode()
        )
        objs.append(_stream(content.encode("ascii")))
    out = bytearray(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for num, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{num} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{off:010d} 00000 n \n".encode() for off in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def make_docx() -> bytes:
    """제목·문단·표(가로 병합, 세로 병합)가 섞인 한국어 DOCX."""
    from docx import Document

    doc = Document()
    doc.add_heading(KOREAN_LINES[0], level=1)
    doc.add_paragraph(KOREAN_LINES[1])
    table = doc.add_table(rows=3, cols=3)
    for col, head in enumerate(["구분", "규모", "출처"]):
        table.cell(0, col).text = head
    table.cell(1, 0).merge(table.cell(2, 0)).text = "데이터"  # 세로 병합: 두 행에 걸친 셀
    table.cell(1, 1).text = "12,000건"
    table.cell(1, 2).text = "문헌"
    table.cell(2, 1).merge(table.cell(2, 2)).text = "300건 자체 실험"  # 가로 병합
    doc.add_paragraph(KOREAN_LINES[3])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_zip(entries: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in entries.items():
            zf.writestr(name, data)
    return buf.getvalue()


def inflate_docx(extra_bytes: int) -> bytes:
    """압축 폭탄형 DOCX: 본문에 같은 문단을 되풀이해 넣어 압축 해제 크기만 키운다(파일은 작다)."""
    para = "<w:p><w:r><w:t>반복 문단 반복 문단 반복 문단</w:t></w:r></w:p>".encode("utf-8")
    src = zipfile.ZipFile(io.BytesIO(make_docx()))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == "word/document.xml":
                assert b"<w:body>" in data
                data = data.replace(b"<w:body>", b"<w:body>" + para * (extra_bytes // len(para)), 1)
            out.writestr(info.filename, data)
    return buf.getvalue()


def docx_with_paragraphs(paragraphs: list[str]) -> bytes:
    from docx import Document

    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def varied_lines(n_chars: int, width: int = 80) -> list[str]:
    """글자 수가 n_chars 근처인 서로 다른 한국어 줄들(압축비가 폭탄처럼 높지 않게)."""
    lines, total, i = [], 0, 0
    while total < n_chars:
        line = f"{i:05d} 연구 방법 {i * 7919 % 100003} 데이터 {i * 104729 % 99991} 평가 지표 설명 "
        line = (line * (width // len(line) + 1))[:width]
        lines.append(line)
        total += len(line) + 1
        i += 1
    return lines


HWP5_BYTES = (
    b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    + b"\x00" * 504
    + "FileHeader".encode("utf-16-le")
    + b"\x00" * 64
    + b"HWP Document File"
    + b"\x00" * 256
)
HWPX_BYTES = make_zip(
    {"mimetype": b"application/hwp+zip", "Contents/content.hpf": b"<opf/>", "Contents/section0.xml": b"<hs:sec/>"}
)
"""본문이 빈 최소 HWPX(구역 1개, 문단 없음). 받는 HWPX 견본은 `tests/e4/test_upload_hwpx.py`의 `make_hwpx`."""


def reject(filename: str, data: bytes) -> UploadRejected:
    with pytest.raises(UploadRejected) as info:
        extract_plan(filename, data)
    return info.value


# ── txt·md ──────────────────────────────────────────────────────────────────


def test_txt_utf8_korean_preserved_exactly():
    result = extract_plan("계획서.txt", PLAN_MD.encode("utf-8"))
    assert result.kind == "txt"
    assert result.encoding == "utf-8"
    assert result.text == PLAN_BODY
    assert result.lines == len(PLAN_BODY.split("\n"))
    assert result.warnings == ()
    assert parse_plan_upload("계획서.txt", PLAN_MD.encode("utf-8")) == PLAN_BODY


def test_md_kind_and_blank_lines_kept():
    body = "# 연구 목표\n\n\n\n- 항목 1\n- 항목 2"
    result = extract_plan("plan.md", body.encode("utf-8"))
    assert result.kind == "md"
    assert result.text == body  # txt·md는 사용자가 쓴 빈 줄을 줄이지 않는다(줄 번호 보존)
    assert result.lines == 6


def test_txt_cp949_fallback_with_warning():
    result = extract_plan("plan.txt", PLAN_CP949.encode("cp949"))
    assert result.encoding == "cp949"
    assert result.text == PLAN_CP949
    assert any("CP949" in w for w in result.warnings)


def test_txt_bom_crlf_and_nfd_normalized():
    nfd = unicodedata.normalize("NFD", "한국어 계획서\r\n둘째 줄\t끝   \r\n")
    assert nfd != unicodedata.normalize("NFC", nfd)
    result = extract_plan("plan.txt", b"\xef\xbb\xbf" + nfd.encode("utf-8"))
    assert result.encoding == "utf-8-sig"
    assert result.text == "한국어 계획서\n둘째 줄\t끝"
    assert result.lines == 2


def test_txt_utf16_bom():
    result = extract_plan("plan.txt", "연구 목표\n방법".encode("utf-16"))
    assert result.encoding == "utf-16"
    assert result.text == "연구 목표\n방법"


def test_txt_undecodable_bytes():
    mostly_ok = ("plan line " * 60).encode("ascii") + b"\xff"
    result = extract_plan("plan.txt", mostly_ok)
    assert result.text.endswith("\ufffd")
    assert any("�" in w for w in result.warnings)
    assert reject("plan.txt", b"\xff" * 200).status_code == 422  # 거의 전부 깨지면 텍스트로 보지 않는다


def test_text_matches_plan_document_lines():
    result = extract_plan("plan.md", PLAN_MD.encode("utf-8"))
    plan = PlanDocument.from_text(result.text, session_id="s1")
    assert len(plan.lines) == result.lines
    assert plan.line(1) == "# 연구계획서 (예시) — 전해액 이온전도도 예측 대리모델"


# ── pdf ─────────────────────────────────────────────────────────────────────


def test_pdf_korean_extracted():
    data = make_pdf([KOREAN_LINES[:2], KOREAN_LINES[2:]])
    result = extract_plan("계획서.pdf", data)
    assert result.kind == "pdf"
    assert result.pages == 2
    assert result.encoding is None
    assert result.text.split("\n") == KOREAN_LINES
    assert result.lines == 4
    assert result.warnings == ()


def test_pdf_fixture_plan_round_trip():
    lines = [ln for ln in PLAN_BODY.split("\n") if ln.strip()]
    result = extract_plan("plan.pdf", make_pdf([lines]))
    assert result.text.split("\n") == lines


def test_pdf_blank_page_warns():
    result = extract_plan("plan.pdf", make_pdf([KOREAN_LINES, [], KOREAN_LINES[:1]]))
    assert result.pages == 3
    assert any("텍스트가 없는 쪽: 2" in w and "OCR" in w for w in result.warnings)
    assert result.text.split("\n") == KOREAN_LINES + KOREAN_LINES[:1]


def test_pdf_without_text_rejected():
    err = reject("scan.pdf", make_pdf([[], []]))
    assert err.status_code == 422
    assert err.message == upload.SCANNED_PDF_MESSAGE
    assert "텍스트가 없는 PDF(스캔본)는 처리할 수 없습니다" in err.message


def test_pdf_encrypted_rejected():
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf([KOREAN_LINES]))))
    writer.encrypt(user_password="secret-pw", algorithm="RC4-128")
    buf = io.BytesIO()
    writer.write(buf)
    err = reject("locked.pdf", buf.getvalue())
    assert err.status_code == 422
    assert "암호" in err.message


def test_pdf_corrupt_rejected():
    err = reject("broken.pdf", b"%PDF-1.7\n" + b"garbage " * 50)
    assert err.status_code == 422


# ── docx ────────────────────────────────────────────────────────────────────


def test_docx_paragraphs_and_tables_in_order():
    result = extract_plan("계획서.docx", make_docx())
    assert result.kind == "docx"
    assert result.pages is None
    assert result.text.split("\n") == [
        KOREAN_LINES[0],
        KOREAN_LINES[1],
        "구분 | 규모 | 출처",
        "데이터 | 12,000건 | 문헌",
        "300건 자체 실험",  # 세로 병합의 연속 셀은 비어 있어 되풀이하지 않는다
        KOREAN_LINES[3],
    ]
    assert result.warnings == ()


def test_docx_corrupt_rejected():
    err = reject("plan.docx", make_zip({"word/document.xml": b"<not-xml", "[Content_Types].xml": b"<x/>"}))
    assert err.status_code == 422
    assert reject("plan.docx", b"PK\x03\x04 truncated").status_code == 422
    assert reject("plan.docx", "그냥 텍스트".encode("utf-8")).status_code == 422


# ── 형식 판별 ───────────────────────────────────────────────────────────────


def test_content_wins_over_extension_with_warning():
    pdf_as_txt = extract_plan("plan.txt", make_pdf([KOREAN_LINES]))
    assert pdf_as_txt.kind == "pdf"
    assert any("확장자(.txt)" in w and "PDF" in w for w in pdf_as_txt.warnings)
    docx_no_ext = extract_plan("plan", make_docx())
    assert docx_no_ext.kind == "docx"
    assert any("확장자(없음)" in w for w in docx_no_ext.warnings)


def test_pdf_marker_inside_text_stays_text():
    body = "PDF 머리 표식은 %PDF-1.7 형태다.\n본문 둘째 줄"
    result = extract_plan("notes.md", body.encode("utf-8"))
    assert result.kind == "md"
    assert result.text == body


def test_control_chars_removed():
    result = extract_plan("plan.txt", "첫 줄\x07끝\x0c둘째 줄\x1b".encode("utf-8"))  # 벨·폼피드·ESC
    assert result.text == "첫 줄끝\n둘째 줄"


def test_filename_path_stripped():
    assert extract_plan("C:\\fakepath\\내 계획서.md", b"x").filename == "내 계획서.md"
    assert extract_plan("../../etc/plan.txt", b"x").filename == "plan.txt"


# ── 거부: 크기·HWP·미지원 ─────────────────────────────────────────────────────


def padded_10mb() -> bytes:
    """정확히 10MB. 줄 끝 공백은 정리에서 빠지므로 글자 상한과 따로 바이트 상한만 잰다."""
    head = "연구 목표\n".encode("utf-8")
    return head + b" " * (MAX_UPLOAD_BYTES - len(head))


def test_size_limit_boundary():
    exact = padded_10mb()
    assert len(exact) == MAX_UPLOAD_BYTES == 10 * 1024 * 1024
    ok = extract_plan("big.txt", exact)
    assert ok.size_bytes == MAX_UPLOAD_BYTES and ok.text == "연구 목표"
    err = reject("big.txt", exact + b"x")
    assert err.status_code == 413
    assert "10MB" in err.message


# ── 상한: SEC-1 S-03 업로드 증폭 ───────────────────────────────────────────────


def test_char_limit_boundary():
    assert MAX_PLAN_CHARS == 50_000
    assert len(extract_plan("plan.txt", ("가" * MAX_PLAN_CHARS).encode("utf-8")).text) == MAX_PLAN_CHARS
    err = reject("plan.txt", ("가" * (MAX_PLAN_CHARS + 1)).encode("utf-8"))
    assert err.status_code == 413
    assert err.message == upload.TOO_MANY_CHARS_MESSAGE and "50,000자" in err.message


def test_char_limit_pdf_and_docx():
    lines = varied_lines(MAX_PLAN_CHARS + 5_000)
    pdf = make_pdf([lines[i : i + 40] for i in range(0, len(lines), 40)])
    assert reject("long.pdf", pdf).message == upload.TOO_MANY_CHARS_MESSAGE
    assert reject("long.docx", docx_with_paragraphs(lines)).message == upload.TOO_MANY_CHARS_MESSAGE
    under = varied_lines(MAX_PLAN_CHARS - 5_000)
    assert extract_plan("ok.docx", docx_with_paragraphs(under)).lines == len(under)


def test_docx_bomb_rejected_fast(monkeypatch):
    bomb = inflate_docx(MAX_ZIP_UNCOMPRESSED + 5 * 1024 * 1024)  # 압축 해제 25MB
    assert len(bomb) < 500 * 1024  # 파일 자체는 작다
    opened: list[str] = []

    def forbid_payload(self, name, *args, **kwargs):
        opened.append(str(name))
        raise AssertionError("zip 폭탄의 압축 내용을 열면 안 된다")

    monkeypatch.setattr(zipfile.ZipFile, "open", forbid_payload)
    err = reject("bomb.docx", bomb)
    assert err.status_code == 413 and err.message == upload.ZIP_BOMB_MESSAGE
    # 확장자를 바꿔도 zip 목록 단계에서 막힌다
    assert reject("bomb.txt", bomb).message == upload.ZIP_BOMB_MESSAGE
    assert opened == [], "풀지 않고 목록만 보고 막아야 한다"


def test_docx_high_ratio_entry_rejected():
    bomb = inflate_docx(5 * 1024 * 1024)  # 합계 상한(20MB) 아래지만 한 항목 압축비가 100배를 넘는다
    info = max(zipfile.ZipFile(io.BytesIO(bomb)).infolist(), key=lambda i: i.file_size)
    assert info.file_size < MAX_ZIP_UNCOMPRESSED and info.file_size > 100 * info.compress_size
    assert reject("ratio.docx", bomb).message == upload.ZIP_BOMB_MESSAGE


def test_zip_entry_count_rejected():
    entries = {f"word/media/x{i}.xml": b"<x/>" for i in range(upload.MAX_ZIP_ENTRIES)}
    entries["word/document.xml"] = b"<w:document/>"
    assert reject("many.docx", make_zip(entries)).message == upload.ZIP_BOMB_MESSAGE


def test_pdf_page_limit_boundary():
    assert MAX_PDF_PAGES == 200
    ok = extract_plan("200.pdf", make_pdf([[f"{i}쪽 본문"] for i in range(MAX_PDF_PAGES)]))
    assert ok.pages == MAX_PDF_PAGES and ok.lines == MAX_PDF_PAGES
    err = reject("201.pdf", make_pdf([[f"{i}쪽 본문"] for i in range(MAX_PDF_PAGES + 1)]))
    assert err.status_code == 413 and err.message == upload.TOO_MANY_PAGES_MESSAGE


def test_cooperative_deadline():
    for name, data in (("p.pdf", make_pdf([KOREAN_LINES])), ("p.docx", make_docx()), ("p.txt", b"plan")):
        with pytest.raises(UploadRejected) as info:
            extract_plan(name, data, deadline_s=-1)
        assert info.value.status_code == 413 and info.value.message == upload.TIMEOUT_MESSAGE


def test_isolated_extraction_runs_in_worker(monkeypatch):
    def broken(*args, **kwargs):
        raise RuntimeError("부모 프로세스에서 추출하면 안 된다")

    monkeypatch.setattr(upload, "_extract_pdf", broken)
    monkeypatch.setattr(upload, "_extract_docx", broken)
    pdf = extract_plan_isolated("p.pdf", make_pdf([KOREAN_LINES]))
    assert pdf.kind == "pdf" and pdf.text.split("\n") == KOREAN_LINES
    assert extract_plan_isolated("p.docx", make_docx()).kind == "docx"
    with pytest.raises(RuntimeError):
        extract_plan("p.pdf", make_pdf([KOREAN_LINES]))  # 대조: 같은 프로세스 경로는 막힌 함수를 쓴다


def test_isolated_rejections_pass_through():
    # 작업자 안에서 난 거부도 같은 상태·문구로 돌아온다
    err = pytest.raises(UploadRejected, extract_plan_isolated, "201.pdf", make_pdf([["a"]] * (MAX_PDF_PAGES + 1))).value
    assert err.status_code == 413 and err.message == upload.TOO_MANY_PAGES_MESSAGE
    # HWP·zip 폭탄은 작업자를 띄우기 전에 거부된다
    assert pytest.raises(UploadRejected, extract_plan_isolated, "a.hwp", HWP5_BYTES).value.status_code == 415
    bomb = inflate_docx(MAX_ZIP_UNCOMPRESSED + 1024 * 1024)
    assert pytest.raises(UploadRejected, extract_plan_isolated, "b.docx", bomb).value.message == upload.ZIP_BOMB_MESSAGE


def test_isolated_timeout_kills_worker(monkeypatch):
    # 시험이 만든 작업자를 확실히 미완료로 둔다. 부모가 상한 뒤 kill/wait 했는지 직접 검사한다.
    monkeypatch.setattr(upload, "_worker_command", lambda: [
        sys.executable, "-c", "import sys,time; sys.stdin.buffer.read(); time.sleep(60)",
    ])
    original = upload.subprocess.Popen
    workers, killed = [], []

    def spawn(*args, **kwargs):
        proc = original(*args, **kwargs)
        workers.append(proc)
        kill = proc.kill

        def record_kill():
            killed.append(proc.pid)
            return kill()

        monkeypatch.setattr(proc, "kill", record_kill)
        return proc

    monkeypatch.setattr(upload.subprocess, "Popen", spawn)
    with pytest.raises(UploadRejected) as info:
        extract_plan_isolated("p.pdf", make_pdf([KOREAN_LINES]), timeout_s=0.05)
    assert info.value.status_code == 413 and info.value.message == upload.TIMEOUT_MESSAGE
    assert len(workers) == 1 and killed == [workers[0].pid]
    assert workers[0].poll() is not None, "시간 초과 뒤 시험 작업자가 남아 있다"


def test_concurrency_limit_503(monkeypatch):
    slots = threading.BoundedSemaphore(1)
    monkeypatch.setattr(upload, "_EXTRACT_SLOTS", slots)
    monkeypatch.setattr(upload, "EXTRACT_QUEUE_WAIT_S", 0.05)
    assert slots.acquire(timeout=1)
    try:
        with pytest.raises(UploadRejected) as info:
            extract_plan_isolated("p.txt", b"plan")
        assert info.value.status_code == 503 and info.value.message == upload.BUSY_MESSAGE
    finally:
        slots.release()
    assert extract_plan_isolated("p.txt", b"plan").text == "plan"  # 자리가 나면 다시 받는다


def test_worker_env_has_no_secrets(monkeypatch):
    monkeypatch.setenv("NEUMANN_FAKE_API_KEY", "x")
    monkeypatch.setenv("FAKE_SERVICE_TOKEN", "x")
    monkeypatch.setenv("NEUMANN_FAKE_SALT", "x")
    env = upload._worker_env()
    assert not {"NEUMANN_FAKE_API_KEY", "FAKE_SERVICE_TOKEN", "NEUMANN_FAKE_SALT", "NEUMANN_PSEUDONYM_SALT"} & set(env)
    assert "OPENAI_API_KEY" not in env
    assert "PATH" in env


def test_worker_denies_disk_writes(tmp_path):
    target = tmp_path / "leak.txt"
    src = str(Path(upload.__file__).resolve().parents[2])
    code = (
        f"import sys; sys.path.insert(0, {src!r}); from neumann.api.upload import _deny_disk_writes; "
        f"_deny_disk_writes(); open(sys.executable, 'rb').close(); print('read-ok', flush=True); open({str(target)!r}, 'w')"
    )
    proc = subprocess.run([sys.executable, "-I", "-B", "-c", code], capture_output=True, timeout=60)
    assert proc.returncode != 0
    assert b"read-ok" in proc.stdout  # 읽기는 된다
    assert b"PermissionError" in proc.stderr
    assert not target.exists()


@pytest.mark.parametrize(
    ("filename", "data"),
    [
        ("계획서.hwp", HWP5_BYTES),
        ("계획서.hwp", b"anything"),  # 확장자만으로도 거부
        ("양식.hwt", HWP5_BYTES),
        ("계획서.hml", "<?xml version=\"1.0\"?><HWPML/>".encode("utf-8")),
        ("disguised.pdf", HWP5_BYTES),  # 매직바이트로 거부
        ("disguised.hwpx", HWP5_BYTES),
        ("old.hwp", b"HWP Document File V3.00 \x1a\x01\x02\x03\x04\x05"),
        ("old.txt", b"HWP Document File V3.00 \x1a\x01\x02\x03\x04\x05"),
    ],
    ids=["hwp5", "hwp-ext-only", "hwt", "hml", "hwp5-as-pdf", "hwp5-as-hwpx", "hwp3", "hwp3-as-txt"],
)
def test_hwp_rejected_415(filename, data):
    """옛 한글 바이너리(.hwp·.hwt, HWP 3·5)와 .hml은 읽지 않는다(대표 지시). HWPX·PDF로 저장하라고 안내한다."""
    err = reject(filename, data)
    assert err.status_code == 415
    assert err.message == HWP_MESSAGE == "HWP는 한글에서 HWPX 또는 PDF로 저장해 올려 주세요"


@pytest.mark.parametrize(
    ("filename", "data", "status", "needle"),
    [
        ("계획서.hwpx", HWPX_BYTES, 422, "HWPX를 읽을 수 없습니다"),  # 구역 XML이 깨졌다(hs: 접두어 선언 없음)
        ("disguised.docx", HWPX_BYTES, 422, "HWPX를 읽을 수 없습니다"),  # 내용 기준 HWPX
        ("계획서.hwpx", "텍스트".encode("utf-8"), 422, "HWPX 파일이 아니거나 손상"),
        ("계획서.hwpx", make_zip({"a.txt": b"x"}), 422, "HWPX 파일이 아니거나 손상"),
    ],
    ids=["hwpx-empty", "hwpx-as-docx", "hwpx-ext-text", "hwpx-ext-plain-zip"],
)
def test_hwpx_not_hwp_message(filename, data, status, needle):
    """HWPX는 더 이상 415(HWP 안내)가 아니다. 빈 본문·가짜 HWPX는 422 안내."""
    err = reject(filename, data)
    assert (err.status_code, err.message != HWP_MESSAGE) == (status, True)
    assert needle in err.message


@pytest.mark.parametrize(
    ("filename", "data", "needle"),
    [
        ("old.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600, "DOCX"),
        ("image.png", b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR", "txt·md·pdf·docx"),
        ("slides.pptx", make_zip({"ppt/presentation.xml": b"<p/>"}), "txt·md·pdf·docx"),
        ("plan.rtf", b"{\\rtf1 hello}", "txt·md·pdf·docx"),
        ("binary.txt", b"abc\x00\x01\x02def", "txt·md·pdf·docx"),
    ],
    ids=["doc", "png", "pptx", "rtf", "binary-as-txt"],
)
def test_unsupported_rejected_415(filename, data, needle):
    err = reject(filename, data)
    assert err.status_code == 415
    assert needle in err.message


def test_empty_rejected():
    assert reject("plan.txt", b"").status_code == 422
    assert reject("plan.txt", b"  \n\n\t\n").status_code == 422


# ── HTTP: POST /upload/plan ─────────────────────────────────────────────────


def _app() -> FastAPI:
    app = FastAPI()
    app.include_router(router)
    return app


@pytest.fixture
def client() -> TestClient:
    return TestClient(_app())


def post(client: TestClient, filename: str, data: bytes, mime: str = "application/octet-stream"):
    return client.post("/upload/plan", files={"file": (filename, data, mime)})


def test_route_registered():
    assert any(getattr(r, "path", None) == "/upload/plan" and "POST" in r.methods for r in router.routes)
    app = FastAPI()
    app.include_router(router)
    body = app.openapi()["paths"]["/upload/plan"]["post"]["requestBody"]
    assert "multipart/form-data" in body["content"]


def test_main_app_wires_upload_router():
    """main.py(PM)가 선택 라우터로 이 모듈을 붙였는지. 붙지 않았으면 /upload/plan이 404다."""
    from neumann.api.main import ROUTER_STATE, app

    assert ROUTER_STATE.get("neumann.api.upload") == "ok"
    res = post(TestClient(app), "계획서.md", "# 연구 목표\n본문".encode("utf-8"), "text/markdown")
    assert res.status_code == 200, res.text
    assert res.json()["lines"] == 2


def test_api_txt(client):
    res = post(client, "연구계획서.md", PLAN_MD.encode("utf-8"), "text/markdown")
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["filename"] == "연구계획서.md"
    assert body["kind"] == "md"
    assert body["text"] == PLAN_BODY
    assert body["lines"] == len(PLAN_BODY.split("\n"))
    assert body["chars"] == len(PLAN_BODY)
    assert body["chars_no_space"] == len("".join(PLAN_BODY.split()))
    assert body["paragraphs"] == sum(1 for line in PLAN_BODY.split("\n") if line.strip())
    assert body["size_bytes"] == len(PLAN_MD.encode("utf-8"))
    assert body["warnings"] == []


def test_api_pdf_and_docx(client):
    res = post(client, "plan.pdf", make_pdf([KOREAN_LINES, []]), "application/pdf")
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["kind"] == "pdf" and body["pages"] == 2 and body["lines"] == 4
    assert body["text"].split("\n") == KOREAN_LINES
    assert any("텍스트가 없는 쪽: 2" in w for w in body["warnings"])

    res = post(client, "plan.docx", make_docx())
    assert res.status_code == 200, res.text
    assert res.json()["kind"] == "docx"
    assert "데이터 | 12,000건 | 문헌" in res.json()["text"].split("\n")


def test_api_cp949_warning(client):
    res = post(client, "plan.txt", PLAN_CP949.encode("cp949"), "text/plain")
    assert res.status_code == 200
    assert res.json()["encoding"] == "cp949"
    assert res.json()["text"] == PLAN_CP949
    assert any("CP949" in w for w in res.json()["warnings"])


def test_api_hwp_415(client):
    for name, data in (("계획서.hwp", HWP5_BYTES), ("양식.hwt", b"x"), ("fake.pdf", HWP5_BYTES)):
        res = post(client, name, data)
        assert res.status_code == 415, name
        assert res.json()["detail"] == HWP_MESSAGE


def test_api_size_limits(client):
    exact = padded_10mb()
    ok = post(client, "big.txt", exact, "text/plain")
    assert ok.status_code == 200
    assert ok.json()["size_bytes"] == MAX_UPLOAD_BYTES
    # 1바이트 초과: 본문 길이 선검사는 통과하고 스트림 수집 단계에서 막힌다
    over = post(client, "big.txt", exact + b"x", "text/plain")
    assert over.status_code == 413
    # 크게 초과: Content-Length 선검사에서 막힌다
    way_over = post(client, "big.txt", exact + b"x" * 200_000, "text/plain")
    assert way_over.status_code == 413
    assert "10MB" in way_over.json()["detail"]


def test_api_amplification_limits(client):
    """SEC-1 S-03 재현형: 작은 DOCX 폭탄·과다 쪽수·과다 글자는 빨리 413으로 끝나고 응답이 작다."""
    cases = [
        ("bomb.docx", inflate_docx(MAX_ZIP_UNCOMPRESSED + 17 * 1024 * 1024), upload.ZIP_BOMB_MESSAGE),  # 압축 해제 37MB
        ("201.pdf", make_pdf([["본문"]] * (MAX_PDF_PAGES + 1)), upload.TOO_MANY_PAGES_MESSAGE),
        ("long.docx", docx_with_paragraphs(varied_lines(MAX_PLAN_CHARS + 1_000)), upload.TOO_MANY_CHARS_MESSAGE),
        ("long.txt", ("가" * (MAX_PLAN_CHARS + 1)).encode("utf-8"), upload.TOO_MANY_CHARS_MESSAGE),
    ]
    for name, data, message in cases:
        start = time.process_time()
        res = post(client, name, data)
        assert res.status_code == 413, (name, res.text)
        assert res.json()["detail"] == message
        assert len(res.content) < 1024
        # 부모의 요청 처리 계산 예산은 유지한다. 자식 추출 상한은 timeout 시험이 별도로 검사한다.
        assert time.process_time() - start < 10.0, name


def test_api_request_shape_errors(client):
    assert client.post("/upload/plan", content=b"hello", headers={"content-type": "text/plain"}).status_code == 415
    only_field = client.post(
        "/upload/plan",
        content=b"--XyZ\r\nContent-Disposition: form-data; name=\"note\"\r\n\r\nhi\r\n--XyZ--\r\n",
        headers={"content-type": "multipart/form-data; boundary=XyZ"},
    )
    assert only_field.status_code == 422
    assert "file" in only_field.json()["detail"]
    mixed_case = client.post(
        "/upload/plan",
        content=(
            "--XyZ\r\nContent-Disposition: form-data; name=\"file\"; filename=\"계획.txt\"\r\n"
            "Content-Type: text/plain\r\n\r\n연구 목표\r\n--XyZ--\r\n"
        ).encode("utf-8"),
        headers={"content-type": "Multipart/Form-Data; boundary=XyZ"},
    )
    assert mixed_case.status_code == 200, mixed_case.text
    assert mixed_case.json()["filename"] == "계획.txt"
    assert mixed_case.json()["text"] == "연구 목표"
    two = client.post(
        "/upload/plan",
        files=[("file", ("a.txt", b"a", "text/plain")), ("file", ("b.txt", b"b", "text/plain"))],
    )
    assert two.status_code == 400
    empty = post(client, "empty.txt", b"", "text/plain")
    assert empty.status_code == 422


def test_api_overlong_boundary_400():
    """경계 문자열이 python-multipart 상한(256자)을 넘으면 500이 아니라 400 + 사용자 문구(검증 지적)."""
    client = TestClient(_app(), raise_server_exceptions=False)

    def send(boundary: str):
        body = (
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"plan.txt\"\r\n"
            f"Content-Type: text/plain\r\n\r\n연구 목표\r\n--{boundary}--\r\n"
        ).encode("utf-8")
        return client.post("/upload/plan", content=body, headers={"content-type": f"multipart/form-data; boundary={boundary}"})

    assert send("b" * 256).status_code == 200  # 상한 안은 그대로 받는다
    for n in (257, 4000):
        res = send("b" * n)
        assert res.status_code == 400, (n, res.text)
        assert res.json()["detail"] == "multipart 본문을 해석할 수 없습니다"


# ── 디스크 미저장 ────────────────────────────────────────────────────────────


@pytest.fixture
def disk_writes(monkeypatch) -> list[str]:
    """임시 파일 생성·쓰기 모드 열기를 기록하고 막는다. Starlette UploadFile 생성도 기록한다."""
    attempts: list[str] = []

    def forbid(name):
        def _raise(*args, **kwargs):
            attempts.append(name)
            raise AssertionError(f"디스크 쓰기 시도: {name}")

        return _raise

    for name in ("TemporaryFile", "NamedTemporaryFile", "mkstemp", "mkdtemp"):
        monkeypatch.setattr(tempfile, name, forbid(f"tempfile.{name}"))

    real_open = builtins.open

    def guarded_open(file, mode="r", *args, **kwargs):
        if isinstance(mode, str) and any(c in mode for c in "wax+"):
            attempts.append(f"open({mode})")
            raise AssertionError("디스크 쓰기 시도: open")
        return real_open(file, mode, *args, **kwargs)

    monkeypatch.setattr(builtins, "open", guarded_open)

    real_os_open = os.open
    write_flags = os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_APPEND

    def guarded_os_open(path, flags, *args, **kwargs):
        if flags & write_flags:
            attempts.append("os.open")
            raise AssertionError("디스크 쓰기 시도: os.open")
        return real_os_open(path, flags, *args, **kwargs)

    monkeypatch.setattr(os, "open", guarded_os_open)

    real_spooled = starlette.formparsers.SpooledTemporaryFile

    def spooled(*args, **kwargs):
        attempts.append("starlette.SpooledTemporaryFile")
        return real_spooled(*args, **kwargs)

    monkeypatch.setattr(starlette.formparsers, "SpooledTemporaryFile", spooled)
    return attempts


def test_api_never_touches_disk(disk_writes):
    app = FastAPI()
    app.include_router(router)
    client = TestClient(app)
    # 약 2.6MB(Starlette 스풀 한도 1MB 초과). 줄 끝 공백은 정리에서 빠져 글자 상한 안에 든다
    big = ("가나다라 계획서 본문 줄" + " " * 2000 + "\n").encode("utf-8") * 1300
    assert len(big) > 2 * 1024 * 1024
    for name, data in (("big.txt", big), ("plan.pdf", make_pdf([KOREAN_LINES])), ("plan.docx", make_docx())):
        res = client.post("/upload/plan", files={"file": (name, data, "application/octet-stream")})
        assert res.status_code == 200, (name, res.text)
    assert disk_writes == []


def test_disk_guard_catches_default_upload_path(disk_writes):
    """대조군: FastAPI 기본 UploadFile은 1MB가 넘으면 임시 파일로 내려 쓴다 → 감시가 실제로 잡는다."""
    app = FastAPI()

    @app.post("/default")
    async def default_upload(file: UploadFile = File(...)):  # noqa: B008
        return {"size": len(await file.read())}

    client = TestClient(app, raise_server_exceptions=False)
    res = client.post("/default", files={"file": ("big.txt", b"x" * (2 * 1024 * 1024), "text/plain")})
    assert res.status_code != 200
    assert "starlette.SpooledTemporaryFile" in disk_writes
    assert any(a.startswith(("tempfile.", "os.open")) for a in disk_writes)
