"""E4-L1f 화면 파일 업로드 ↔ ``POST /upload/plan`` 연결 검사(기본 pytest, 브라우저 없음).

1) API 쪽: 화면이 읽는 필드(``text``·``filename``·``kind``·``pages``·``encoding``·``warnings``, 오류 ``detail``)가
   실제 응답에 있는지 main 앱(라우터 연결 포함)에 txt·md·docx·pdf·hwp를 보내 본다.
2) 화면 쪽(index.html 정적 검사): ``readFile``이 ``/upload/plan``을 부르는지, "준비 중" 문구가 없는지,
   파일명·경고·서버 문구가 ``esc()``를 거쳐서만 HTML에 들어가는지, 드롭존 안내·``accept``가 스펙대로인지.
   실제 브라우저 동작(업로드 → 본문 반영, 415 문구, 중복 제출 차단, 정적 판)은 ``test_webui_upload_ui.py``(Playwright).

LLM을 부르지 않는다. 그래도 mock provider를 강제한다(사용자 환경 변수가 openai여도).
"""

from __future__ import annotations

import io
import re
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from neumann.api import main
from neumann.api.upload import HWP_MESSAGE, MAX_UPLOAD_BYTES, PlanUploadResponse
from tests.e4.test_upload import HWP5_BYTES, HWPX_BYTES, make_pdf

ROOT = Path(__file__).resolve().parents[2]
INDEX = ROOT / "src" / "neumann" / "webui" / "index.html"

UI_FIELDS = {"text", "filename", "kind", "pages", "encoding", "warnings"}
DROP_HINT = "TXT · MD · PDF · DOCX · 최대 10 MB · 정리 뒤 50,000자 · HWP는 PDF·DOCX로 저장"
LIVE_ONLY = "PDF·DOCX는 라이브 서버에서만 읽습니다 — 본문을 직접 입력에 붙여넣기"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PLAN_LINES = [
    "연구계획서 — 전해액 이온전도도 예측 대리모델",
    "1. 목표: GNN으로 이온전도도(mS/cm)를 예측한다.",
    "2. 데이터: 문헌 12,000건, 자체 실험 300건",
]


def make_docx(lines: list[str] = PLAN_LINES) -> bytes:
    """python-docx로 제목 + 문단 DOCX를 만든다(파일 커밋 없음)."""
    from docx import Document

    doc = Document()
    doc.add_heading(lines[0], level=1)
    for line in lines[1:]:
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture(autouse=True)
def _mock_only(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("NEUMANN_LLM_PROVIDER", "mock")
    monkeypatch.delenv("NEUMANN_LIVE_TESTS", raising=False)


@pytest.fixture(scope="module")
def client() -> TestClient:
    return TestClient(main.app)


def _post(client: TestClient, name: str, data: bytes):
    return client.post("/upload/plan", files={"file": (name, data)})


def _ui_shape(body: dict) -> None:
    assert UI_FIELDS <= body.keys(), body.keys()
    assert isinstance(body["text"], str) and body["text"]
    assert isinstance(body["filename"], str) and body["filename"]
    assert body["kind"] in {"txt", "md", "pdf", "docx"}
    assert body["pages"] is None or isinstance(body["pages"], int)
    assert body["encoding"] is None or isinstance(body["encoding"], str)
    assert isinstance(body["warnings"], list) and all(isinstance(w, str) for w in body["warnings"])
    assert isinstance(body["size_bytes"], int) and body["size_bytes"] > 0


# ── 1 API 응답 형태(화면이 읽는 필드) ─────────────────────────────────────────


def test_upload_router_is_wired_into_main_app():
    assert main.ROUTER_STATE.get("neumann.api.upload") == "ok"
    assert "post" in main.app.openapi()["paths"]["/upload/plan"]  # FastAPI 0.141은 포함 라우터를 감싸 두므로 스키마로 본다


def test_txt_utf8_response_shape(client):
    body = "\n".join(PLAN_LINES)
    r = _post(client, "계획서.txt", body.encode("utf-8"))
    assert r.status_code == 200, r.text
    j = r.json()
    _ui_shape(j)
    assert (j["kind"], j["filename"], j["pages"], j["encoding"], j["warnings"]) == ("txt", "계획서.txt", None, "utf-8", [])
    assert j["text"] == body


def test_md_cp949_response_carries_encoding_and_warning(client):
    body = "# 연구계획서 - 한글\n둘째 줄"
    r = _post(client, "계획서.md", body.encode("cp949"))
    assert r.status_code == 200, r.text
    j = r.json()
    _ui_shape(j)
    assert (j["kind"], j["encoding"], j["pages"]) == ("md", "cp949", None)
    assert j["text"] == body
    assert any("CP949" in w for w in j["warnings"])  # 화면이 파일 카드 아래에 보여 줄 경고


def test_docx_response_shape(client):
    r = _post(client, "계획서.docx", make_docx())
    assert r.status_code == 200, r.text
    j = r.json()
    _ui_shape(j)
    assert (j["kind"], j["pages"], j["encoding"]) == ("docx", None, None)
    assert j["text"].splitlines() == PLAN_LINES


def test_pdf_response_has_pages_and_empty_page_warning(client):
    r = _post(client, "계획서.pdf", make_pdf([PLAN_LINES, []]))
    assert r.status_code == 200, r.text
    j = r.json()
    _ui_shape(j)
    assert (j["kind"], j["pages"], j["encoding"]) == ("pdf", 2, None)
    assert j["text"].splitlines() == PLAN_LINES
    assert any("텍스트가 없는 쪽" in w for w in j["warnings"])


@pytest.mark.parametrize(("name", "data"), [("계획서.hwp", HWP5_BYTES), ("계획서.hwpx", HWPX_BYTES)])
def test_hwp_is_415_with_server_message(client, name, data):
    r = _post(client, name, data)
    assert r.status_code == 415
    assert r.json() == {"detail": HWP_MESSAGE}  # 화면은 이 detail을 그대로 S.inErr에 쓴다


@pytest.mark.parametrize(
    ("name", "size", "status"),
    [("빈.txt", 0, 422), ("큰.txt", MAX_UPLOAD_BYTES + 1, 413), ("표.xlsx", -1, 415)],
    ids=["empty-422", "over-10MB-413", "xlsx-415"],
)
def test_errors_carry_string_detail(client, name, size, status):
    data = b"PK\x03\x04junk" if size < 0 else b"a" * size
    r = _post(client, name, data)
    assert r.status_code == status
    detail = r.json()["detail"]
    assert isinstance(detail, str) and detail


def test_filename_comes_back_verbatim_so_ui_must_escape(client):
    name = "<img src=x onerror=alert(1)>.txt"  # 따옴표는 multipart에서 %22로 바뀌므로 넣지 않는다
    r = _post(client, name, "본문".encode())
    assert r.status_code == 200
    assert r.json()["filename"] == name  # 서버는 파일명을 바꾸지 않는다 → 화면이 esc()로 넣어야 한다


# ── 2 index.html 정적 검사 ────────────────────────────────────────────────────


def _html() -> str:
    return INDEX.read_text(encoding="utf-8")


def _function(src: str, name: str) -> str:
    """``function name(...) { ... }`` 본문을 중괄호 짝으로 잘라 낸다(문자열 안 중괄호는 이 파일 범위에서 짝이 맞는다)."""
    m = re.search(r"function " + re.escape(name) + r"\s*\([^)]*\)\s*\{", src)
    assert m, f"{name} 함수가 없다"
    depth, i = 0, m.end() - 1
    while True:
        c = src[i]
        depth += c == "{"
        depth -= c == "}"
        i += 1
        if depth == 0:
            return src[m.start() : i]


def _const(src: str, name: str) -> str:
    m = re.search(r"var " + name + r" = '([^']*)';", src)
    assert m, f"{name} 상수가 없다"
    return m.group(1)


def test_readfile_posts_formdata_to_upload_plan():
    src = _html()
    fn = _function(src, "readFile")
    assert "fetch('upload/plan', { method: 'POST', body: fd })" in fn
    assert "fd.append('file', file, name)" in fn
    assert "FileReader" not in fn  # 서버 판은 브라우저에서 읽지 않는다(정적 판 readLocal에서만)
    assert "FileReader" in _function(src, "readLocal")


def test_pending_text_is_gone():
    src = _html()
    assert "PDF · DOCX 준비 중" not in src
    assert "업로드는 준비 중" not in src


def test_drop_hint_and_accept():
    src = _html()
    assert _const(src, "UP_HINT") == DROP_HINT
    accept = _const(src, "UP_ACCEPT").split(",")
    for want in (".txt", ".md", ".pdf", ".docx", "text/plain", "text/markdown", "application/pdf", DOCX_MIME):
        assert want in accept
    drop = _function(src, "dropHtml")
    assert "UP_HINT" in drop and "accept=\"' + UP_ACCEPT + '\"" in drop
    assert "area = dropHtml() +" in _function(src, "renderInput")


def test_filename_warnings_and_errors_only_through_esc():
    src = _html()
    render_input = _function(src, "renderInput")
    assert "esc(S.filename)" in render_input and "esc(S.fileMeta)" in render_input
    # S.filename·S.fileMeta·S.inErr가 esc() 없이 HTML 문자열에 붙는 곳이 파일 전체에 없어야 한다
    for field in ("filename", "fileMeta", "inErr"):
        assert not re.search(r"\+\s*S\." + field + r"\b", src), field
    drop = _function(src, "dropHtml")
    assert "esc(b)" in drop and not re.search(r"\+\s*b\s*\+", drop)
    warn = _function(src, "fileWarnHtml")
    assert "'<li>' + esc(x) + '</li>'" in warn
    assert "fileWarnHtml() +" in render_input  # 경고는 파일 카드 바로 아래
    for name in ("readFile", "readLocal", "upDone", "upDetail"):
        assert "innerHTML" not in _function(src, name), name
    for name in ("UP_HINT", "UP_ACCEPT", "UP_LIVE_ONLY", "UP_OFFLINE"):
        assert not re.search(r"[<>\"&]", _const(src, name)), name


def test_error_detail_shown_verbatim_and_network_message():
    src = _html()
    detail = _function(src, "upDetail")
    assert "typeof d === 'string' && d) return d;" in detail  # 서버 detail 그대로
    fn = _function(src, "readFile")
    assert "S.inErr = upDetail(j, x.status)" in fn
    assert _const(src, "UP_OFFLINE") == "서버에 연결하지 못함"
    assert "readLocal(file, name, UP_OFFLINE)" in fn  # fetch 실패(네트워크) 경로


def test_busy_state_blocks_duplicate_submit():
    src = _html()
    fn = _function(src, "readFile")
    assert fn.index("if (!file || UP.busy) return;") < fn.index("fetch(")
    assert "UP.busy = name || '파일'; render();" in fn
    drop = _function(src, "dropHtml")
    assert "'읽는 중…'" in drop and "' disabled'" in drop and "aria-busy" in drop


def test_static_mode_reads_only_md_txt_locally():
    src = _html()
    assert _const(src, "UP_LIVE_ONLY") == LIVE_ONLY
    fn = _function(src, "readFile")
    assert "r.status === 404" in fn and "UP.api = 'none'" in fn
    local = _function(src, "readLocal")
    assert "ext === 'pdf' || ext === 'docx'" in local and "UP_LIVE_ONLY" in local
    assert "/^(md|txt|markdown)$/" in local
    assert "브라우저에서 UTF-8로 읽음" in local  # 폴백으로 읽었다고 파일 카드에 남긴다


def test_ui_reads_only_fields_the_api_returns():
    fn = _function(_html(), "readFile")
    used = set(re.findall(r"\bj\.(\w+)", fn))
    assert UI_FIELDS <= used, used
    assert used <= set(PlanUploadResponse.model_fields) | {"detail"}, used - set(PlanUploadResponse.model_fields)
