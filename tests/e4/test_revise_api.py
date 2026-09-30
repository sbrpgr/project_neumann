"""E3-L2r API 검사: POST /premortem/revise · /premortem/revise/assemble이 서빙 관문(차단·속도 제한·상한·대기열·시간 상한)을
그대로 거치고, 계약 JSON·md·docx를 돌려주며, 관문 없는 요청은 받지 않는다. 실제 파이프라인·OpenAI 없음(mock provider)."""

from __future__ import annotations

import asyncio
import io
import json
import time
from pathlib import Path
from typing import Any

import httpx
import pytest
from fastapi import FastAPI

from neumann.analyze import revise as revise_mod
from neumann.analyze.revise import validate_revision
from neumann.api import revise as revise_api
from neumann.api import serving
from scripts.serve_fake_app import fake_result, integrate
from tests.e3.revise_fixtures import make_store
from tests.fixtures.loader import load_fixtures, plan_text

PLAN = plan_text("plan.md")


@pytest.fixture(autouse=True)
def _records(monkeypatch):
    monkeypatch.setattr(revise_mod, "default_record_store", lambda: make_store())
    monkeypatch.setenv("NEUMANN_LLM_PROVIDER", "mock")
    monkeypatch.delenv("NEUMANN_REVISE_TIMEOUT_S", raising=False)


def make(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, **cfg: Any) -> tuple[serving.Serving, Any]:
    base: dict[str, Any] = dict(max_concurrent=2, queue_max=4, rate_per_min=0, cache_enabled=False,
                                cache_dir=tmp_path / "results", request_timeout_s=30.0, avg_run_s=0.1,
                                block_file=tmp_path / "serving_block.flag")
    base.update(cfg)
    srv = serving.Serving(serving.ServingConfig(**base))
    app = integrate(srv, lambda plan_text: fake_result(plan_text), monkeypatch.setattr)
    assert revise_api.install(app, srv=srv)
    return srv, app


def client(app: Any, ip: str = "127.0.0.1") -> httpx.AsyncClient:
    return httpx.AsyncClient(transport=httpx.ASGITransport(app=app, client=(ip, 5000)), base_url="http://t", timeout=30)


def result_json() -> dict[str, Any]:
    return load_fixtures().premortem_result.model_dump(mode="json")


def run(coro):
    return asyncio.run(coro)


# ── /premortem/revise ─────────────────────────────────────────────────────


def test_revise_happy_path_contract_and_gate_release(tmp_path, monkeypatch):
    srv, app = make(tmp_path, monkeypatch)

    async def go() -> None:
        async with client(app) as c:
            r = await c.post("/premortem/revise", json={"result": result_json(), "plan_text": PLAN})
            assert r.status_code == 200, r.text
            body = r.json()
            assert validate_revision(body) == []
            assert body["generator"] == "mock" and len(body["revisions"]) == 2 and body["cost"]["llm_calls"] == 2
            assert "revision_sig" in body  # 서명 모듈(E4-L2f)이 없으면 null
            assert r.headers["x-neumann-ticket"] and r.headers["cache-control"] == "no-store"
            assert "Traceback" not in r.text and "C:\\\\" not in r.text
            one = await c.post("/premortem/revise", json={"result": result_json(), "plan_text": PLAN, "card_ids": ["card-fx-leak"]})
            assert one.status_code == 200 and [x["card_id"] for x in one.json()["revisions"]] == ["card-fx-leak"]
            assert one.json()["cards_requested"] == ["card-fx-leak"]
    run(go())
    assert srv.gate.active == 0 and srv.gate.waiting == 0
    assert srv.counters["requests"] == 2


def test_revise_rejects_mismatched_plan_and_invalid_result(tmp_path, monkeypatch):
    _srv, app = make(tmp_path, monkeypatch)

    async def go() -> None:
        async with client(app) as c:
            r = await c.post("/premortem/revise", json={"result": result_json(), "plan_text": PLAN + "\n바뀐 줄"})
            assert r.status_code == 422 and r.json()["error_code"] == "plan_mismatch"
            r = await c.post("/premortem/revise", json={"result": {"session_id": "x", "plan_id": "y", "bogus": "SECRET_VALUE_9f3"},
                                                        "plan_text": PLAN})
            assert r.status_code == 422 and r.json()["error_code"] == "invalid_request"
            assert "SECRET_VALUE_9f3" not in r.text  # 입력 값을 되돌려 보내지 않는다(위치·종류만)
            r = await c.post("/premortem/revise", json={"plan_text": PLAN})
            assert r.status_code == 422
    run(go())


def test_revise_goes_through_block_switch_rate_limit_and_size_cap(tmp_path, monkeypatch):
    srv, app = make(tmp_path, monkeypatch, rate_per_min=1, max_plan_chars=len(PLAN) + 10)

    async def go() -> None:
        async with client(app) as c:
            ok = await c.post("/premortem/revise", json={"result": result_json(), "plan_text": PLAN})
            assert ok.status_code == 200
            limited = await c.post("/premortem/revise", json={"result": result_json(), "plan_text": PLAN})
            assert limited.status_code == 429 and limited.json()["error_code"] == "rate_limited" and limited.headers["retry-after"]
        async with client(app, ip="10.0.0.7") as c:
            big = await c.post("/premortem/revise", json={"result": result_json(), "plan_text": PLAN + "x" * 100})
            assert big.status_code == 413 and big.json()["error_code"] == "too_large"
            (tmp_path / "serving_block.flag").write_text("1")
            blocked = await c.post("/premortem/revise", json={"result": result_json(), "plan_text": PLAN})
            assert blocked.status_code == 503 and blocked.json()["error_code"] == "blocked"
    run(go())
    assert srv.counters["rate_429"] >= 1 and srv.counters["too_large_413"] >= 1 and srv.counters["blocked_503"] >= 1


def test_revise_second_line_admission_when_middleware_skipped_reservation(tmp_path, monkeypatch):
    """같은 계획서 분석이 캐시에 있으면 미들웨어는 자리를 잡지 않는다. 그래도 핸들러가 입장 검사를 다시 거친다(차단 → 503)."""
    srv, app = make(tmp_path, monkeypatch, cache_enabled=True)
    pid = serving.plan_key(PLAN)
    srv.cache.put(pid, {**fake_result(PLAN), "plan_id": pid, "status": "ok"})
    assert srv.cache.has(pid)

    async def go() -> None:
        async with client(app) as c:
            ok = await c.post("/premortem/revise", json={"result": result_json(), "plan_text": PLAN})
            assert ok.status_code == 200
            (tmp_path / "serving_block.flag").write_text("1")
            blocked = await c.post("/premortem/revise", json={"result": result_json(), "plan_text": PLAN})
            assert blocked.status_code == 503 and blocked.json()["error_code"] == "blocked"
    run(go())
    assert srv.gate.active == 0


def test_revise_not_gated_is_refused(tmp_path, monkeypatch):
    app = FastAPI()
    app.include_router(revise_api.router)  # 서빙 층 없이 붙인 경우(관문 없는 경로) → 받지 않는다

    async def go() -> None:
        async with client(app) as c:
            r = await c.post("/premortem/revise", json={"result": result_json(), "plan_text": PLAN})
            assert r.status_code == 503 and r.json()["error_code"] == "unavailable"
    run(go())
    assert revise_api.install(app) is False  # serving.install이 없으면 붙이지 않는다


def test_revise_timeout_is_user_message(tmp_path, monkeypatch):
    srv, app = make(tmp_path, monkeypatch)
    monkeypatch.setenv("NEUMANN_REVISE_TIMEOUT_S", "0.3")

    def slow(req, **kw):
        time.sleep(1.0)
        return {}

    monkeypatch.setattr(revise_api, "run_revision", slow)

    async def go() -> None:
        async with client(app) as c:
            r = await c.post("/premortem/revise", json={"result": result_json(), "plan_text": PLAN})
            assert r.status_code == 504 and r.json()["error_code"] == "timeout"
            assert "0초" in r.json()["message"] and "Traceback" not in r.text
    run(go())
    assert srv.gate.active == 0


def test_revise_internal_error_is_user_message(tmp_path, monkeypatch):
    _srv, app = make(tmp_path, monkeypatch)

    def boom(req, **kw):
        raise RuntimeError("C:\\Users\\someone\\secret\\revise_internal.py")

    monkeypatch.setattr(revise_api, "run_revision", boom)

    async def go() -> None:
        async with client(app) as c:
            r = await c.post("/premortem/revise", json={"result": result_json(), "plan_text": PLAN})
            assert r.status_code == 500 and r.json()["error_code"] == "internal"
            assert "someone" not in r.text and "RuntimeError" not in r.text
    run(go())


# ── /premortem/revise/assemble ───────────────────────────────────────────


async def _bundle(c: httpx.AsyncClient) -> dict[str, Any]:
    r = await c.post("/premortem/revise", json={"result": result_json(), "plan_text": PLAN})
    assert r.status_code == 200
    return r.json()


def _decisions(bundle: dict[str, Any]) -> list[dict[str, Any]]:
    edits = [e for rev in bundle["revisions"] for e in rev["edits"]]
    return [{"edit_id": edits[0]["edit_id"], "decision": "채택"},
            {"edit_id": edits[1]["edit_id"], "decision": "수정", "revised_text": "조성 그룹 단위 분할로 바꾸고 중복은 분할 전에 제거한다."}]


def test_assemble_json_md_docx_and_polish(tmp_path, monkeypatch):
    from docx import Document

    from neumann.analyze.assemble import validate_revised_plan

    _srv, app = make(tmp_path, monkeypatch)

    async def go() -> None:
        async with client(app) as c:
            bundle = await _bundle(c)
            body = {"plan_text": PLAN, "revision": bundle, "decisions": _decisions(bundle), "result": result_json()}
            r = await c.post("/premortem/revise/assemble", json=body)
            assert r.status_code == 200, r.text
            out = r.json()
            assert validate_revised_plan(out) == []
            assert out["stats"]["applied"] == 2 and out["polish"] == {**out["polish"], "requested": False, "applied": False}
            assert set(out["markdown"]) == {"clean", "footnoted", "history"} and out["label"].startswith("Neumann 수정 제안")
            assert out["docx_available"] is True and "revised_plan_sig" in out
            assert out["markdown"]["clean"] == out["revised_text"]
            md = await c.post("/premortem/revise/assemble", json={**body, "format": "md"})
            assert md.status_code == 200 and md.headers["content-type"].startswith("text/markdown")
            assert "[^1]" in md.text and "## 수정 이력" in md.text
            dx = await c.post("/premortem/revise/assemble", json={**body, "format": "docx", "title": "수정본 시험"})
            assert dx.status_code == 200 and dx.headers["content-type"] == revise_api.DOCX_MEDIA
            assert dx.content[:2] == b"PK" and dx.headers["x-neumann-changes"] == "2"
            doc = Document(io.BytesIO(dx.content))
            assert doc.paragraphs[0].text == "수정본 시험" and doc.tables
            pol = await c.post("/premortem/revise/assemble", json={**body, "polish": True})
            assert pol.status_code == 200 and pol.json()["polish"]["applied"] is True and pol.json()["polish"]["generator"] == "mock"
    run(go())


def test_assemble_rejects_bad_revision_and_mismatch(tmp_path, monkeypatch):
    _srv, app = make(tmp_path, monkeypatch)

    async def go() -> None:
        async with client(app) as c:
            bundle = await _bundle(c)
            r = await c.post("/premortem/revise/assemble", json={"plan_text": PLAN, "revision": {"version": "x"}, "decisions": []})
            assert r.status_code == 422 and r.json()["error_code"] == "invalid_request"
            r = await c.post("/premortem/revise/assemble", json={"plan_text": PLAN + "\n다른 계획서", "revision": bundle, "decisions": []})
            assert r.status_code == 422 and r.json()["error_code"] == "plan_mismatch"
            other = {**result_json(), "plan_id": "0" * 64, "plan": None}
            r = await c.post("/premortem/revise/assemble", json={"plan_text": PLAN, "revision": bundle, "decisions": [], "result": other})
            assert r.status_code == 422 and r.json()["error_code"] == "plan_mismatch"
            r = await c.post("/premortem/revise/assemble", json={"plan_text": PLAN, "revision": bundle, "decisions": [], "format": "pdf"})
            assert r.status_code == 422
    run(go())


def test_assemble_conflict_listed_via_api(tmp_path, monkeypatch):
    _srv, app = make(tmp_path, monkeypatch)

    async def go() -> None:
        async with client(app) as c:
            bundle = await _bundle(c)
            e0 = bundle["revisions"][0]["edits"][0]
            bundle["revisions"][0]["edits"].append({**e0, "edit_id": e0["edit_id"] + "b", "proposed_text": "다른 안"})
            decisions = [{"edit_id": e0["edit_id"], "decision": "adopt"}, {"edit_id": e0["edit_id"] + "b", "decision": "adopt"}]
            r = await c.post("/premortem/revise/assemble", json={"plan_text": PLAN, "revision": bundle, "decisions": decisions})
            assert r.status_code == 200
            out = r.json()
            assert out["stats"]["conflicts"] == 1 and out["conflicts"][0]["kind"] == "same_line" and out["stats"]["applied"] == 0
    run(go())


def test_main_app_registers_revise_paths_as_protected():
    from neumann.api import main as api_main

    srv = api_main.app.state.serving
    assert srv.kind_for("/premortem/revise") == "analysis" and srv.kind_for("/premortem/revise/assemble") == "analysis"
    paths = set(api_main.app.openapi()["paths"])
    assert {"/premortem/revise", "/premortem/revise/assemble"} <= paths


def test_plan_id_of_matches_pipeline_normalization():
    fx = load_fixtures()
    assert revise_api.plan_id_of(PLAN) == fx.premortem_result.plan_id
    assert revise_api.plan_id_of(PLAN + "\n연락처 010-1234-5678") != revise_api.plan_id_of(PLAN)
    assert json.dumps(revise_api.sign_payload("revision", {"a": 1.0}) or None) in ("null", json.dumps(revise_api.sign_payload("revision", {"a": 1})))
