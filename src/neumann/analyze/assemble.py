"""통합 단계(E3-L2r, 대표 정의 "완전한 유기체 = 근거 기반으로 고쳐진 연구계획서 한 부").

    assembled = assemble_revised_plan(plan_text, revision, decisions)      # 채택·직접 수정한 안을 줄 단위로 합친다(코드만)
    assembled = polish_revised_plan(assembled, llm_call, effort="medium")  # 선택: 문장 다듬기 LLM 1회 + 코드 게이트
    md = render_markdown(assembled, evidence)                              # (a) 깨끗한 원고 (b) 각주 판 (c) 수정 이력 + 근거 부록
    data = build_docx(assembled, evidence, model=..., generated_at=...)    # .docx(python-docx, 미주·수정 이력 표·근거 부록)

규칙
- 적용하는 안은 연구자가 **채택**(제안 문안 그대로)하거나 **수정**(연구자 문안 `revised_text`)한 것뿐이다. 기각·미결정은 적용하지
  않는다. 결정이 없는 안은 그대로 두고 `undecided`에 센다.
- 같은 줄을 여러 안이 바꾸면(같은 줄 replace 둘 이상) **충돌 목록**으로 돌려주고 자동으로 고르지 않는다(그 줄은 원문 유지).
  안의 `current_text`가 계획서 줄과 다르면(계획서가 바뀜) `stale_line` 충돌이다. `insert_after`는 같은 줄에 여럿이어도
  순서대로 넣는다(충돌 아님).
- 변경 문장마다 근거 id(수정안 이유의 excerpt id)와 카드 id를 붙인다. 자리표시(`[확인 필요: …]`)는 따로 목록에 담는다.
- 다듬기(polish)는 바뀐 줄의 흐름만 고친다. 코드 게이트: 줄 수·번호 같음, 바꾸지 않은 줄은 글자 그대로, 바뀐 줄은 수치 집합·
  자리표시 집합이 같고 길이 비율이 0.5~2.0이며 개인정보가 없다. 하나라도 어기면 다듬기 전체를 버리고 통합본을 쓴다.
- 연구자 문안·메모는 이메일·ORCID를 가린다. 리뷰어 신원 필드는 없다(근거 발췌는 원문 오프셋 인용뿐).
"""

from __future__ import annotations

import io
import json
import logging
import re
from collections.abc import Callable, Mapping, Sequence
from copy import copy
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Any

from neumann.analyze import gate as gate_mod
from neumann.analyze.revise import (CONTROL_RE, PLACEHOLDER_RE, PROPOSED_LABEL, UNSAFE_MARKUP_RE, contains_identity,
                                    placeholders, unsupported_facts, written_numbers)
from neumann.models import PlanDocument, contains_pii, redact_pii

log = logging.getLogger(__name__)

ASSEMBLE_VERSION = "revised_plan@v1"
POLISH_TASK = "polish_plan"
POLISH_PROMPT_VERSION = "polish_plan@v1"
POLISH_GATE_VERSION = "polish-diff-gate@v1"
DECISIONS = ("adopt", "modify", "reject")
DECISION_ALIASES = {"채택": "adopt", "수정": "modify", "기각": "reject", "adopt": "adopt", "modify": "modify", "reject": "reject",
                    "accept": "adopt", "edit": "modify", "rejected": "reject", "adopted": "adopt"}
DECISION_KO = {"adopt": "채택", "modify": "수정", "reject": "기각"}
MAX_REVISED_CHARS = 2000
LABEL_PREFIX = "Neumann 수정 제안"

POLISH_INSTRUCTIONS = """\
You smooth the wording of a revised research plan so that the changed lines read naturally with their neighbours.

Input (JSON): lines of the plan with a line number and a flag `changed`. Lines with changed=false must be returned
verbatim (byte for byte). Lines with changed=true may be reworded for flow only.

Hard rules (the system checks them and discards the whole output if any is broken):
1. Return every line with the same number, in the same order, and nothing else.
2. Do not add, remove or alter any number, claim, result, dataset, method, resource or commitment. Rewording only.
3. Keep every placeholder of the form [확인 필요: ...] exactly as it is.
4. Do not add quotation marks, ids, or comments. Keep the language of each line (Korean or English).
"""


# ── 입력 정리 ─────────────────────────────────────────────────────────────


@dataclass
class EditRef:
    edit_id: str
    card_id: str
    plan_line: int
    kind: str
    proposed_text: str
    current_text: str
    excerpt_ids: list[str] = field(default_factory=list)
    rationale: str = ""


@dataclass
class DecisionRef:
    edit_id: str
    decision: str
    revised_text: str | None = None
    note: str | None = None
    decided_at: str | None = None


def normalize_decision(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    return DECISION_ALIASES.get(value.strip().lower(), DECISION_ALIASES.get(value.strip()))


def collect_edits(revision: Mapping[str, Any] | Sequence[Mapping[str, Any]]) -> dict[str, EditRef]:
    """수정 권고 묶음(`revisions[].edits[]`) 또는 수정안 목록 → edit_id별 참조."""
    items: list[tuple[str, Mapping[str, Any]]] = []
    if isinstance(revision, Mapping):
        for rev in revision.get("revisions", []) or []:
            if isinstance(rev, Mapping):
                for e in rev.get("edits", []) or []:
                    if isinstance(e, Mapping):
                        items.append((str(rev.get("card_id", "")), e))
    else:
        for e in revision:
            if isinstance(e, Mapping):
                items.append((str(e.get("card_id", "")), e))
    out: dict[str, EditRef] = {}
    for card_id, e in items:
        eid = str(e.get("edit_id") or "")
        no = e.get("plan_line")
        if not eid:
            continue
        if isinstance(no, bool) or not isinstance(no, int):
            raise ValueError("edit plan_line must be an integer")
        rat = e.get("rationale") if isinstance(e.get("rationale"), Mapping) else {}
        ids = [x for x in (rat.get("excerpt_ids") or e.get("excerpt_ids") or []) if isinstance(x, str)]
        out.setdefault(eid, EditRef(
            edit_id=eid, card_id=card_id or str(e.get("card_id", "")), plan_line=no,
            kind=str(e.get("kind") or "replace"), proposed_text=str(e.get("proposed_text") or ""),
            current_text=str(e.get("current_text") or ""), excerpt_ids=ids, rationale=str(rat.get("text") or ""),
        ))
    return out


def collect_decisions(decisions: Sequence[Mapping[str, Any]] | None) -> tuple[dict[str, DecisionRef], list[str]]:
    """결정 목록 → edit_id별 마지막 결정. 형식 오류는 사유 목록으로."""
    out: dict[str, DecisionRef] = {}
    problems: list[str] = []
    for i, raw in enumerate(decisions or ()):
        if not isinstance(raw, Mapping):
            problems.append(f"결정 {i}: 객체가 아니다")
            continue
        eid = raw.get("edit_id")
        dec = normalize_decision(raw.get("decision"))
        if not isinstance(eid, str) or not eid:
            problems.append(f"결정 {i}: edit_id가 없다")
            continue
        if dec is None:
            problems.append(f"결정 {i}({eid}): decision은 채택·수정·기각(adopt·modify·reject) 중 하나다")
            continue
        text = raw.get("revised_text")
        text = redact_pii(" ".join(CONTROL_RE.sub("", str(text)).split())) if isinstance(text, str) and text.strip() else None
        if text is not None and len(text) > MAX_REVISED_CHARS:
            problems.append(f"결정 {i}({eid}): revised_text가 {MAX_REVISED_CHARS}자를 넘는다")
            continue
        note = raw.get("note")
        note = redact_pii(CONTROL_RE.sub("", str(note)))[:2000] if isinstance(note, str) and note.strip() else None
        at = raw.get("decided_at")
        out[eid] = DecisionRef(eid, dec, text, note, str(at) if isinstance(at, str) else None)
    return out, problems


# ── 통합 ─────────────────────────────────────────────────────────────────


def _plan_of(plan_text: str | PlanDocument) -> PlanDocument:
    if isinstance(plan_text, PlanDocument):
        return plan_text
    try:
        from neumann.pipeline import mask_extra_pii

        masked, _ = mask_extra_pii(plan_text)
    except Exception:  # noqa: BLE001
        masked = plan_text
    return PlanDocument.from_text(masked, "assemble")


def assemble_revised_plan(
    plan_text: str | PlanDocument,
    revision: Mapping[str, Any] | Sequence[Mapping[str, Any]],
    decisions: Sequence[Mapping[str, Any]] | None = None,
    *,
    result: Any = None,
    regate: bool = False,
) -> dict[str, Any]:
    """채택·수정된 안을 줄 단위로 합친다(LLM 없음). 계약 `contracts/revised_plan.schema.json`의 본체.

    반환: revised_text, lines(번호·본문·changed·edit_id), changes(줄 범위·원문·수정문·카드·근거 id·결정), conflicts(미해결),
    placeholders, stats, notices, polish(요청 전이면 {"requested": false, "applied": false}).
    """
    plan = _plan_of(plan_text)
    edits = collect_edits(revision)
    decided, problems = collect_decisions(decisions)
    notices: list[str] = list(problems)
    rev_plan_id = revision.get("plan_id") if isinstance(revision, Mapping) else None
    if isinstance(rev_plan_id, str) and rev_plan_id and rev_plan_id != plan.plan_id:
        notices.append("수정 권고의 plan_id가 계획서 본문과 다르다 — 줄 번호가 어긋날 수 있다(current_text 대조로 걸러진다)")
    original = {ln.no: ln.text for ln in plan.lines}
    if regate:
        _regate_adoptions(plan, revision, edits, decided, result)

    applied: dict[str, tuple[EditRef, DecisionRef, str]] = {}  # edit_id → (안, 결정, 적용 문안)
    skipped: list[dict[str, Any]] = []
    conflicts: list[dict[str, Any]] = []
    for eid, d in decided.items():
        e = edits.get(eid)
        if e is None:
            skipped.append({"edit_id": eid, "reason": "unknown_edit", "detail": "수정 권고에 없는 edit_id"})
            continue
        if d.decision == "reject":
            continue
        if d.decision == "modify" and not d.revised_text:
            skipped.append({"edit_id": eid, "reason": "modify_without_text", "detail": "수정 결정에 revised_text가 없다"})
            continue
        if e.plan_line not in original:
            conflicts.append({"kind": "unknown_line", "plan_line": e.plan_line, "edit_ids": [eid], "card_ids": [e.card_id],
                              "detail": "계획서에 없는 줄"})
            continue
        if e.current_text != original[e.plan_line]:
            conflicts.append({"kind": "stale_line", "plan_line": e.plan_line, "edit_ids": [eid], "card_ids": [e.card_id],
                              "detail": "안의 current_text가 계획서 줄과 다르다(계획서가 바뀌었다)",
                              "current_text": original[e.plan_line], "expected_text": e.current_text})
            continue
        text = d.revised_text if d.decision == "modify" else e.proposed_text
        text = " ".join(text.split())
        if not text:
            skipped.append({"edit_id": eid, "reason": "empty_text", "detail": "적용할 문안이 비었다"})
            continue
        applied[eid] = (e, d, text)

    # 같은 줄 replace 둘 이상 → 충돌(자동 선택 없음, 그 줄은 원문 유지)
    by_line_replace: dict[int, list[str]] = {}
    for eid, (e, _d, _t) in applied.items():
        if e.kind == "replace":
            by_line_replace.setdefault(e.plan_line, []).append(eid)
    for no, ids in sorted(by_line_replace.items()):
        if len(ids) > 1:
            conflicts.append({
                "kind": "same_line", "plan_line": no, "edit_ids": ids, "card_ids": [applied[i][0].card_id for i in ids],
                "detail": "같은 줄을 바꾸는 안이 둘 이상이다 — 연구자가 하나를 고르거나 직접 합쳐 수정한다",
                "candidates": [{"edit_id": i, "card_id": applied[i][0].card_id, "text": applied[i][2],
                                "decision": applied[i][1].decision} for i in ids],
                "current_text": original[no],
            })
            for i in ids:
                applied.pop(i, None)

    lines: list[dict[str, Any]] = []
    changes: list[dict[str, Any]] = []
    order = list(applied)  # 결정 순서(입력 순서) 유지
    for no in sorted(original):
        text = original[no]
        rep = next((i for i in order if applied[i][0].kind == "replace" and applied[i][0].plan_line == no), None)
        if rep is not None:
            e, d, new_text = applied[rep]
            new_no = len(lines) + 1
            lines.append({"no": new_no, "orig_no": no, "text": new_text, "changed": True, "edit_id": rep, "card_id": e.card_id})
            changes.append(_change(e, d, no, text, new_text, new_no))
        else:
            lines.append({"no": len(lines) + 1, "orig_no": no, "text": text, "changed": False})
        for ins in [i for i in order if applied[i][0].kind == "insert_after" and applied[i][0].plan_line == no]:
            e, d, new_text = applied[ins]
            new_no = len(lines) + 1
            lines.append({"no": new_no, "orig_no": None, "text": new_text, "changed": True, "edit_id": ins, "card_id": e.card_id})
            changes.append(_change(e, d, no, "", new_text, new_no))

    holders: list[dict[str, Any]] = []
    for ch in changes:
        for ph in placeholders(ch["new_text"]):
            holders.append({"text": ph, "edit_id": ch["edit_id"], "card_id": ch["card_id"], "line": ch["new_range"][0]})
    undecided = [eid for eid in edits if eid not in decided]
    rejected = [eid for eid, d in decided.items() if d.decision == "reject"]
    stats = {
        "edits_total": len(edits), "applied": len(changes), "adopted": sum(1 for c in changes if c["decision"] == "adopt"),
        "modified": sum(1 for c in changes if c["decision"] == "modify"), "rejected": len(rejected),
        "undecided": len(undecided), "conflicts": len(conflicts), "skipped": len(skipped), "placeholders": len(holders),
        "lines_original": len(original), "lines_revised": len(lines),
    }
    if conflicts:
        notices.append(f"미해결 충돌 {len(conflicts)}건 — 해당 줄은 원문을 유지했다. 화면에서 하나를 고르거나 직접 수정한다")
    if holders:
        notices.append(f"자리표시 {len(holders)}곳 — 연구자만 아는 값을 채워야 한다")
    return {
        "version": ASSEMBLE_VERSION,
        "plan_id": plan.plan_id,
        "revised_plan_id": PlanDocument.from_text("\n".join(ln["text"] for ln in lines), "assemble").plan_id,
        "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "revised_text": "\n".join(ln["text"] for ln in lines),
        "lines": lines,
        "changes": changes,
        "conflicts": conflicts,
        "placeholders": holders,
        "skipped": skipped,
        "undecided": undecided,
        "rejected": rejected,
        "stats": stats,
        "notices": notices,
        "polish": {"requested": False, "applied": False, "reason": None, "generator": None, "model": None},
    }


def _regate_adoptions(plan: PlanDocument, revision: Any, edits: dict[str, EditRef],
                     decided: dict[str, DecisionRef], result: Any) -> None:
    """되돌려 받은 제안은 새 신뢰 경계다. 채택 문안을 적용하기 전에 다시 검사한다."""
    from neumann.analyze.revise import fabricated_numbers
    from neumann.analyze.gate import Draft, EvidenceIndex
    from neumann.models import Excerpt, PremortemResult

    ev = evidence_lookup(result, revision if isinstance(revision, Mapping) else None)
    pools = {str(r.get("card_id")): set(r.get("evidence_pool", []))
             for r in (revision.get("revisions", []) if isinstance(revision, Mapping) else []) if isinstance(r, Mapping)}
    records = {str(rec["excerpt_id"]): rec
               for rec in (revision.get("records", []) if isinstance(revision, Mapping) else [])
               if isinstance(rec, Mapping) and rec.get("excerpt_id")}
    index = EvidenceIndex(result if isinstance(result, PremortemResult) else PremortemResult.model_validate(result)) if result is not None else None
    for eid, decision in decided.items():
        edit = edits.get(eid)
        if decision.decision != "adopt" or edit is None:
            continue
        text = edit.proposed_text
        if contains_pii(text) or contains_identity(text):
            raise ValueError("proposed_text contains personal information")
        if CONTROL_RE.search(text) or UNSAFE_MARKUP_RE.search(text):
            raise ValueError("proposed_text contains unsafe characters or markup")
        if unsupported_facts(text, plan.text):
            raise ValueError("proposed_text contains unsupported facts")
        if "[확인 필요" in PLACEHOLDER_RE.sub("", text):
            raise ValueError("proposed_text has an invalid placeholder")
        if not edit.excerpt_ids or any(x not in ev or x not in pools.get(edit.card_id, set()) for x in edit.excerpt_ids):
            raise ValueError("proposed_text lacks verified evidence links")
        if index is not None:
            source_card = index.cards.get(edit.card_id)
            # 카드 원래 근거 + 이 카드의 풀에 포함된 추가 기록만 허용한다.
            # records와 evidence_pool은 revision 서명 범위에 함께 포함된다.
            allowed_ids = (set(source_card.evidence) if source_card is not None else set()) | (
                pools.get(edit.card_id, set()) & records.keys())
            if source_card is None or any(x not in allowed_ids for x in edit.excerpt_ids):
                raise ValueError("proposed_text cites another card")
            # 수치 대조용 발췌는 제안별 사본에만 추가한다. 앞선 채택이 뒤 검사에 영향을 주지 않는다.
            proposal_index = copy(index)
            proposal_index.excerpts = dict(index.excerpts)
            for x in edit.excerpt_ids:
                if x not in proposal_index.excerpts and x in records:
                    rec = records[x]
                    proposal_index.excerpts[x] = Excerpt.model_validate({k: v for k, v in rec.items()
                                                                        if k in Excerpt.model_fields})
            unknown = fabricated_numbers(text, Draft("edit", text, tuple(edit.excerpt_ids), (), (edit.plan_line,)), proposal_index)
        else:
            allowed = set([*gate_mod.extract_numbers(plan.text), *written_numbers(plan.text)])
            for x in edit.excerpt_ids:
                allowed.update([*gate_mod.extract_numbers(ev[x]["text"]), *written_numbers(ev[x]["text"])])
            unknown = [n for n in [*gate_mod.extract_numbers(text), *written_numbers(text)] if n not in allowed]
        if unknown:
            raise ValueError("proposed_text contains unsupported numbers")


def _change(e: EditRef, d: DecisionRef, no: int, old_text: str, new_text: str, new_no: int) -> dict[str, Any]:
    return {
        "edit_id": e.edit_id, "card_id": e.card_id, "kind": e.kind, "old_range": [no, no], "old_text": old_text,
        "new_text": new_text, "new_range": [new_no, new_no], "excerpt_ids": [] if d.decision == "modify" else list(e.excerpt_ids),
        "rationale": "" if d.decision == "modify" else e.rationale, "decision": d.decision, "decision_ko": DECISION_KO[d.decision],
        "revised_by": "researcher" if d.decision == "modify" else "proposal", "proposed_label": PROPOSED_LABEL,
        "note": d.note, "decided_at": d.decided_at, "placeholders": placeholders(new_text),
    }


# ── 다듬기(선택, LLM 1회 + 코드 게이트) ──────────────────────────────────


def polish_schema(n_lines: int) -> dict[str, Any]:
    return {
        "type": "object", "additionalProperties": False, "required": ["lines"],
        "properties": {"lines": {"type": "array", "items": {
            "type": "object", "additionalProperties": False, "required": ["no", "text"],
            "properties": {"no": {"type": "integer", "minimum": 1, "maximum": max(n_lines, 1)}, "text": {"type": "string"}},
        }}},
    }


def polish_gate(before: list[dict[str, Any]], after: Any) -> tuple[list[dict[str, Any]] | None, str]:
    """다듬기 출력 검사. 통과면 (새 줄 목록, ""), 아니면 (None, 사유)."""
    if not isinstance(after, Mapping) or not isinstance(after.get("lines"), list):
        return None, "output_shape"
    rows = after["lines"]
    if len(rows) != len(before):
        return None, f"line_count {len(rows)} != {len(before)}"
    out: list[dict[str, Any]] = []
    for old, new in zip(before, rows, strict=True):
        if not isinstance(new, Mapping) or new.get("no") != old["no"] or not isinstance(new.get("text"), str):
            return None, f"line_shape at {old['no']}"
        text = new["text"]
        if not old.get("changed"):
            if text != old["text"]:
                return None, f"unchanged_line_modified at {old['no']}"
            out.append(dict(old))
            continue
        text = " ".join(text.split())
        if not text:
            return None, f"empty_line at {old['no']}"
        if contains_pii(text) or contains_identity(text):
            return None, f"pii at {old['no']}"
        if CONTROL_RE.search(text) or UNSAFE_MARKUP_RE.search(text):
            return None, f"unsafe_text at {old['no']}"
        if sorted([*gate_mod.extract_numbers(text), *written_numbers(text)]) != sorted([*gate_mod.extract_numbers(old["text"]), *written_numbers(old["text"])]):
            return None, f"numbers_changed at {old['no']}"
        if sorted(placeholders(text)) != sorted(placeholders(old["text"])):
            return None, f"placeholders_changed at {old['no']}"
        if any(ch in text for ch in "\"“”「」『』《》«»＂") and not any(ch in old["text"] for ch in "\"“”「」『』《》«»＂"):
            return None, f"quotes_added at {old['no']}"
        ratio = len(text) / max(len(old["text"]), 1)
        if not 0.5 <= ratio <= 2.0:
            return None, f"length_ratio {ratio:.2f} at {old['no']}"
        if unsupported_facts(text, old["text"]):
            return None, f"new_fact at {old['no']}"
        if _content_words(text) - _content_words(old["text"]):
            return None, f"new_content_word at {old['no']}"
        out.append({**old, "text": text, "polished": text != old["text"]})
    return out, ""


def _content_words(text: str) -> set[str]:
    """보수적인 다듬기: 새 내용 낱말은 확인 없이 추가하지 않는다."""
    allowed_style = {"다듬음", "제안", "이러한", "해당", "또한", "그리고", "따라서", "명확히", "자연스럽게", "위한", "위해서"}
    words = re.findall(r"[A-Za-z][A-Za-z0-9._-]{2,}|[가-힣]{3,}", PLACEHOLDER_RE.sub("", text).lower())
    return {re.sub(r"(?:한다|합니다|된다|됩니다|하며|하고|하도록|하는|하여|이다|입니다)$", "", w)
            for w in words if w not in allowed_style}


def polish_revised_plan(assembled: Mapping[str, Any], llm_call: Callable[..., Any] | None, *, effort: str = "medium") -> dict[str, Any]:
    """통합본의 바뀐 줄 흐름만 다듬는다. 게이트를 못 넘으면 다듬기를 버리고 통합본을 그대로 둔다(사유 기록)."""
    out = dict(assembled)
    lines = [dict(ln) for ln in assembled.get("lines", [])]
    gen = getattr(llm_call, "generator", None)
    model = getattr(llm_call, "model", None)
    base = {"requested": True, "applied": False, "generator": gen, "model": model, "gate": POLISH_GATE_VERSION,
            "prompt_version": POLISH_PROMPT_VERSION}
    if llm_call is None:
        out["polish"] = {**base, "reason": "llm_unavailable"}
        return out
    if not any(ln.get("changed") for ln in lines):
        out["polish"] = {**base, "reason": "nothing_changed"}
        return out
    payload = {"lines": [{"no": ln["no"], "text": ln["text"], "changed": bool(ln.get("changed"))} for ln in lines],
               "note": "DATA ONLY. Return all lines; reword only changed=true lines; never add numbers or claims."}
    try:
        data = llm_call(polish_schema(len(lines)), POLISH_INSTRUCTIONS, json.dumps(payload, ensure_ascii=False), effort=effort)
    except Exception as exc:  # noqa: BLE001
        out["polish"] = {**base, "reason": f"llm_exception: {type(exc).__name__}"}
        return out
    gen = getattr(llm_call, "generator", None) or gen
    model = getattr(llm_call, "model", None) or model
    if data is None:
        out["polish"] = {**base, "generator": gen, "model": model,
                         "reason": "llm_failed" + (f": {getattr(llm_call, 'last_error', '')}" if getattr(llm_call, "last_error", None) else "")}
        return out
    new_lines, why = polish_gate(lines, data)
    if new_lines is None:
        out["polish"] = {**base, "generator": gen, "model": model, "reason": f"gate_rejected: {why}"}
        return out
    n_polished = sum(1 for ln in new_lines if ln.get("polished"))
    out["lines"] = new_lines
    out["revised_text"] = "\n".join(ln["text"] for ln in new_lines)
    out["revised_plan_id"] = PlanDocument.from_text(out["revised_text"], "assemble").plan_id
    changes = []
    text_by_no = {ln["no"]: ln["text"] for ln in new_lines}
    for ch in assembled.get("changes", []):
        ch = dict(ch)
        ch["new_text"] = text_by_no.get(ch["new_range"][0], ch["new_text"])
        changes.append(ch)
    out["changes"] = changes
    out["polish"] = {**base, "applied": True, "generator": gen, "model": model, "reason": None, "lines_polished": n_polished}
    return out


# ── 렌더링: 마크다운 3판 · docx ──────────────────────────────────────────


EvidenceLookup = Mapping[str, Mapping[str, Any]]


def evidence_lookup(result: Mapping[str, Any] | Any = None, revision: Mapping[str, Any] | None = None) -> dict[str, dict[str, Any]]:
    """excerpt_id → {text, source_url, source_kind, record_kind, work_id}. 결과 evidence + 권고 records. 신원 필드 없음."""
    out: dict[str, dict[str, Any]] = {}
    res = result.model_dump(mode="json") if hasattr(result, "model_dump") else (dict(result) if isinstance(result, Mapping) else {})
    for ex in res.get("evidence", []) or []:
        if isinstance(ex, Mapping) and ex.get("excerpt_id"):
            out[str(ex["excerpt_id"])] = {"text": _display_excerpt(str(ex.get("text", ""))), "source_url": _display_url(str(ex.get("source_url", ""))),
                                          "source_kind": str(ex.get("source_kind", "review")), "record_kind": "review",
                                          "work_id": None}
    for rec in (revision or {}).get("records", []) or []:
        if isinstance(rec, Mapping) and rec.get("excerpt_id"):
            out.setdefault(str(rec["excerpt_id"]), {"text": _display_excerpt(str(rec.get("text", ""))), "source_url": _display_url(str(rec.get("source_url", ""))),
                                                     "source_kind": str(rec.get("source_kind", "")),
                                                     "record_kind": str(rec.get("record_kind", "")), "work_id": rec.get("work_id")})
    return out


def _display_excerpt(text: str) -> str:
    # 개인정보 포함 입력의 문자열을 원문 인용처럼 렌더링하지 않는다.
    if contains_pii(text) or contains_identity(text):
        return "(개인정보가 포함된 발췌는 표시하지 않음)"
    if CONTROL_RE.search(text) or UNSAFE_MARKUP_RE.search(text):
        return "(안전하지 않은 문자가 포함된 발췌는 표시하지 않음)"
    return text


def _display_url(url: str) -> str:
    return redact_pii(url) if url.startswith(("https://", "http://")) and not re.search(r"[<>\x00-\x20]", url) else ""


KIND_KO = {"review": "심사평", "meta_review": "메타리뷰", "author_response": "저자 답변", "decision": "결정", "post_status": "사후 기록"}


def _label_line(model: str | None, generated_at: str, generator: str | None = None) -> str:
    if generator == "client_submitted_unverified":
        return f"{LABEL_PREFIX} · client_submitted_unverified (입력 출처·생성자 미확인) · 생성 시각 {generated_at}"
    who = "LLM" if generator in (None, "astra") else generator
    return f"{LABEL_PREFIX} · {who} ({model or '모델 미상'}) · 생성 시각 {generated_at}"


def _footnotes(assembled: Mapping[str, Any], ev: EvidenceLookup) -> tuple[dict[str, list[int]], list[dict[str, Any]]]:
    """변경마다 각주 번호 목록과 각주 본문(근거 발췌·링크). 같은 발췌는 한 번호."""
    notes: list[dict[str, Any]] = []
    num_of: dict[str, int] = {}
    per_edit: dict[str, list[int]] = {}
    for ch in assembled.get("changes", []):
        nums: list[int] = []
        for x in ch.get("excerpt_ids", []):
            if x not in num_of:
                e = ev.get(x, {})
                num_of[x] = len(notes) + 1
                notes.append({"n": num_of[x], "excerpt_id": x, "kind": KIND_KO.get(str(e.get("record_kind") or e.get("source_kind") or ""), "근거"),
                              "text": str(e.get("text", "")) or "(발췌 본문 없음 — id만)", "url": str(e.get("source_url", ""))})
            nums.append(num_of[x])
        per_edit[ch["edit_id"]] = nums
    return per_edit, notes


def render_markdown(assembled: Mapping[str, Any], ev: EvidenceLookup, *, model: str | None = None,
                    generator: str | None = None, title: str = "수정된 연구계획서") -> dict[str, str]:
    """(a) clean: 깨끗한 원고 (b) footnoted: 변경 문장 옆 각주 → 근거 발췌·링크 (c) history: 수정 이력 + 근거 부록 + 충돌·자리표시."""
    gen_at = str(assembled.get("generated_at", ""))
    label = _label_line(model, gen_at, generator)
    per_edit, notes = _footnotes(assembled, ev)
    clean = str(assembled.get("revised_text", ""))
    fl: list[str] = []
    for ln in assembled.get("lines", []):
        text = str(ln.get("text", ""))
        if ln.get("changed") and per_edit.get(ln.get("edit_id"), []):
            text += "".join(f"[^{n}]" for n in per_edit[ln["edit_id"]])
        fl.append(text)
    fl += ["", "---", "", f"_{label}_", ""]
    for n in notes:
        url = f" — <{n['url']}>" if n["url"] else ""
        fl.append(f"[^{n['n']}]: {n['kind']} `{n['excerpt_id']}`: {json.dumps(n['text'], ensure_ascii=False)}{url}")
    footnoted = "\n".join(fl).rstrip("\n") + "\n"
    h: list[str] = [f"# {title} — 수정 이력과 근거", "", f"_{label}_", "",
                    f"- 원 계획서 id `{assembled.get('plan_id', '')}` · 수정본 id `{assembled.get('revised_plan_id', '')}`"]
    st = assembled.get("stats", {})
    h.append(f"- 적용 {st.get('applied', 0)}건(채택 {st.get('adopted', 0)} · 수정 {st.get('modified', 0)}) · 기각 {st.get('rejected', 0)} · "
             f"미결정 {st.get('undecided', 0)} · 충돌 {st.get('conflicts', 0)} · 자리표시 {st.get('placeholders', 0)}")
    pol = assembled.get("polish", {})
    h.append(f"- 문장 다듬기: {'적용' if pol.get('applied') else '미적용'}" + (f" ({pol.get('reason')})" if pol.get("reason") else ""))
    h += ["", "## 수정 이력", ""]
    if assembled.get("changes"):
        h += ["| # | 원문 줄 | 종류 | 결정 | 카드 | 원문 | 수정문 | 근거 id |", "|---|---|---|---|---|---|---|---|"]
        for i, ch in enumerate(assembled["changes"], 1):
            h.append(f"| {i} | {ch['old_range'][0]} | {ch['kind']} | {ch['decision_ko']}({ch['revised_by']}) | `{ch['card_id']}` | "
                     f"{_cell(ch['old_text'])} | {_cell(ch['new_text'])} | {', '.join(f'`{x}`' for x in ch['excerpt_ids']) or '-'} |")
    else:
        h.append("적용된 변경이 없다.")
    if assembled.get("conflicts"):
        h += ["", "## 미해결 충돌(자동으로 고르지 않음)", ""]
        for c in assembled["conflicts"]:
            h.append(f"- 줄 {c.get('plan_line')} · {c.get('kind')} · 안 {', '.join(c.get('edit_ids', []))}: {c.get('detail', '')}")
    if assembled.get("placeholders"):
        h += ["", "## 자리표시(연구자가 채울 값)", ""]
        h += [f"- 줄 {p['line']} ({p['edit_id']}): {p['text']}" for p in assembled["placeholders"]]
    h += ["", "## 근거 부록", ""]
    if notes:
        for n in notes:
            h.append(f"{n['n']}. {n['kind']} `{n['excerpt_id']}`" + (f" · <{n['url']}>" if n["url"] else ""))
            h += [f"   > {q}" for q in n["text"].split("\n")]
    else:
        h.append("인용된 근거가 없다.")
    h += ["", "## 한계", "",
          "- 수정문은 제안(근거 아님)이거나 연구자가 직접 쓴 문장이다. 근거 발췌는 비슷한 연구가 실제로 받은 심사 기록을 원문 오프셋으로 자른 것이다.",
          "- 자리표시는 연구자만 아는 값의 자리다. 채우기 전에는 계획서로 제출하지 않는다.",
          "- 리뷰어 신원 정보는 없다."]
    return {"clean": clean, "footnoted": footnoted, "history": "\n".join(h).rstrip("\n") + "\n"}


def _cell(text: Any) -> str:
    return re.sub(r"\s+", " ", redact_pii(str(text))).strip().replace("|", "\\|")[:300]


def build_docx(assembled: Mapping[str, Any], ev: EvidenceLookup, *, model: str | None = None, generator: str | None = None,
               generated_at: str | None = None, title: str = "수정된 연구계획서") -> bytes:
    """python-docx로 .docx를 만든다: 제목·표기 줄·본문(변경 문장 옆 위첨자 미주 번호)·미주·수정 이력 표·근거 부록."""
    from docx import Document
    from docx.shared import Pt

    def xml_safe(value: Any) -> Any:
        if isinstance(value, str):
            return CONTROL_RE.sub("", value)
        if isinstance(value, Mapping):
            return {k: xml_safe(v) for k, v in value.items()}
        if isinstance(value, list):
            return [xml_safe(v) for v in value]
        return value

    assembled, ev = xml_safe(assembled), xml_safe(ev)
    title, model, generator, generated_at = (xml_safe(v) for v in (title, model, generator, generated_at))

    gen_at = generated_at or str(assembled.get("generated_at", ""))
    label = _label_line(model, gen_at, generator)
    per_edit, notes = _footnotes(assembled, ev)
    doc = Document()
    doc.add_heading(title, level=0)
    p = doc.add_paragraph()
    r = p.add_run(label)
    r.italic = True
    r.font.size = Pt(9)
    doc.add_paragraph(f"원 계획서 id {assembled.get('plan_id', '')} · 수정본 id {assembled.get('revised_plan_id', '')}").runs[0].font.size = Pt(8)
    doc.add_heading("본문", level=1)
    for ln in assembled.get("lines", []):
        text = str(ln.get("text", ""))
        para = doc.add_paragraph()
        if not text.strip():
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", text)
        if m and not ln.get("changed"):
            para.style = doc.styles[f"Heading {min(len(m.group(1)) + 1, 4)}"]
            para.add_run(m.group(2))
            continue
        run = para.add_run(text)
        if ln.get("changed"):
            run.bold = True
            for n in per_edit.get(ln.get("edit_id"), []):
                sup = para.add_run(str(n))
                sup.font.superscript = True
    doc.add_heading("미주(변경 문장의 근거)", level=1)
    if notes:
        for n in notes:
            para = doc.add_paragraph(style="List Number")
            para.add_run(f"{n['kind']} {n['excerpt_id']}: ").bold = True
            para.add_run(n["text"])
            if n["url"]:
                para.add_run(f" — {n['url']}").italic = True
    else:
        doc.add_paragraph("인용된 근거가 없다.")
    doc.add_heading("수정 이력", level=1)
    changes = list(assembled.get("changes", []))
    if changes:
        table = doc.add_table(rows=1, cols=7)
        table.style = "Table Grid"
        for cell, head in zip(table.rows[0].cells, ("#", "원문 줄", "종류", "결정", "카드", "원문", "수정문"), strict=True):
            cell.text = head
        for i, ch in enumerate(changes, 1):
            row = table.add_row().cells
            for cell, val in zip(row, (str(i), str(ch["old_range"][0]), ch["kind"], f"{ch['decision_ko']}({ch['revised_by']})",
                                       ch["card_id"], ch["old_text"], ch["new_text"]), strict=True):
                cell.text = redact_pii(str(val))
    else:
        doc.add_paragraph("적용된 변경이 없다.")
    if assembled.get("conflicts"):
        doc.add_heading("미해결 충돌", level=2)
        for c in assembled["conflicts"]:
            doc.add_paragraph(f"줄 {c.get('plan_line')} · {c.get('kind')} · 안 {', '.join(c.get('edit_ids', []))}: {c.get('detail', '')}",
                              style="List Bullet")
    if assembled.get("placeholders"):
        doc.add_heading("자리표시(연구자가 채울 값)", level=2)
        for ph in assembled["placeholders"]:
            doc.add_paragraph(f"줄 {ph['line']} ({ph['edit_id']}): {ph['text']}", style="List Bullet")
    doc.add_heading("근거 부록", level=1)
    if notes:
        for n in notes:
            para = doc.add_paragraph()
            para.add_run(f"{n['n']}. {n['kind']} {n['excerpt_id']}").bold = True
            doc.add_paragraph(n["text"], style="Intense Quote")
            if n["url"]:
                doc.add_paragraph(n["url"]).runs[0].font.size = Pt(8)
    else:
        doc.add_paragraph("인용된 근거가 없다.")
    doc.add_paragraph(label).runs[0].font.size = Pt(8)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def validate_revised_plan(data: Mapping[str, Any]) -> list[str]:
    """계약(contracts/revised_plan.schema.json) 위반 목록."""
    import jsonschema

    from neumann.api.view import SCHEMA_PATH

    schema = json.loads((SCHEMA_PATH.parent / "revised_plan.schema.json").read_text(encoding="utf-8"))
    validator = jsonschema.Draft202012Validator(schema)
    errs = sorted(validator.iter_errors(dict(data)), key=lambda e: list(e.absolute_path))
    return [f"{'/'.join(str(p) for p in e.absolute_path) or '(root)'}: {e.message[:200]}" for e in errs]


__all__ = [
    "ASSEMBLE_VERSION",
    "DECISIONS",
    "DECISION_KO",
    "POLISH_GATE_VERSION",
    "POLISH_INSTRUCTIONS",
    "POLISH_TASK",
    "assemble_revised_plan",
    "build_docx",
    "collect_decisions",
    "collect_edits",
    "evidence_lookup",
    "normalize_decision",
    "polish_gate",
    "polish_revised_plan",
    "polish_schema",
    "render_markdown",
    "validate_revised_plan",
]
