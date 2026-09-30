# E6-pres5 check: markers, verbatim definition, dropped claims, PDF fonts/links, pixel diff vs the pre-pres5 preview.
# usage: pres5_check.py <deck.pptx> <deck.pdf> <new preview dir> <old preview dir>
import os
import sys

import numpy as np
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

deck, pdf, prev_new, prev_old = sys.argv[1:5]
D = ("Neumann의 에이전트는 불완전한 연구계획을, 불완전한 거절 기록들을 모아 설명해 주고, "
     "완전한 기획으로 다시 재탄생시켜 주는 에이전트.")
p = Presentation(deck)


def texts(shapes):
    for s in shapes:
        if s.shape_type == 6:
            yield from texts(s.shapes)
            continue
        if s.has_text_frame:
            yield s.text_frame.text
        if getattr(s, "has_table", False) and s.has_table:
            for r in s.table.rows:
                for c in r.cells:
                    yield c.text_frame.text


slides = []
for i, sl in enumerate(p.slides, 1):
    t = "\n".join(texts(sl.shapes))
    n = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
    slides.append((i, t, n))
print("slides", len(slides))
pats = ["【병합 후 확정】", "【측정 중】", "【v1 확정 전】", "【공개 후 기입】",
        D, "에이전트입니다.", "통계적 결론 없음", "통계적 결론 제한", "시연 예시(선별)", "사후 재실행", "판정 v3",
        "Z3", "Pint", "NetworkX", "산술 일관성", "반복 루프 0", "[확인 필요]", "연구자 코드 실행", "sandbox", "calc",
        "10만 편", "4.9만", "데이터 확장", "데이터 1단계", "왜 커야", "대규모", "학습", "원문 복원", "재측정",
        "astra 합성", "gpt-6-astra", "gpt-6.1-sol", "평가 모델", "07:30", "15편", "[현장 작성"]
print("pattern | slide-text count [slides] | notes count [slides]")
for pat in pats:
    st = [(i, t.count(pat)) for i, t, n in slides if pat in t]
    nt = [(i, n.count(pat)) for i, t, n in slides if pat in n]
    print("%-22s | %d %s | %d %s" % (pat[:22], sum(c for _, c in st), [i for i, _ in st],
                                     sum(c for _, c in nt), [i for i, _ in nt]))
print("markers per slide (slide text):")
for i, t, n in slides:
    row = {m: t.count(m) for m in ("【병합 후 확정】", "【측정 중】", "【v1 확정 전】", "【공개 후 기입】") if t.count(m)}
    if row:
        print("  ", i, row)
r = PdfReader(pdf)
allt = [pg.extract_text() or "" for pg in r.pages]
print("PDF pages", len(r.pages))
for pat in ["【병합 후 확정】", "【측정 중】", "【v1 확정 전】", "【공개 후 기입】", "10만 편", "데이터 확장", "에이전트입니다."]:
    print("PDF %-14s %d %s" % (pat, sum(t.count(pat) for t in allt), [i + 1 for i, t in enumerate(allt) if pat in t]))
fonts = set()
for pg in r.pages:
    res = pg.get("/Resources") or {}
    f = res.get("/Font") or {}
    for k in f:
        fonts.add(str(f[k].get_object().get("/BaseFont")))
print("fonts", sorted(fonts))
links = []
for pg in r.pages:
    for a in pg.get("/Annots") or []:
        a = a.get_object()
        u = a.get("/A", {}).get("/URI") if a.get("/A") else None
        if u:
            links.append(u)
print("links", sorted(set(links)))
# new slide k <- old slide m (16 and 19 were dropped)
old_of = {k: k for k in range(1, 16)}
old_of.update({16: 17, 17: 18})
print("pixel diff vs pre-pres5 (count of |diff|>40), new<-old:")
for k in range(1, len(slides) + 1):
    a = np.asarray(Image.open(os.path.join(prev_new, "slide_%02d.png" % k)).convert("L"), dtype=int)
    b = np.asarray(Image.open(os.path.join(prev_old, "slide_%02d.png" % old_of[k])).convert("L"), dtype=int)
    print("%d<-%d %d" % (k, old_of[k], int((abs(a - b) > 40).sum())), end=" · ")
print()
