# Build the working deck from the kit original. All Korean text comes from content.json (UTF-8).
# usage: build_deck.py <orig.pptx> <content.json> <out_stage1.pptx> <img_dir>
import copy
import json
import os
import re
import sys

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.text.text import _Run
from pptx.util import Emu, Pt

ORIG, CONTENT, OUT, IMGDIR = sys.argv[1:5]
C = json.load(open(CONTENT, encoding="utf-8"))
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
MARK_FG, MARK_BG, MARK = C["mark_fg"], C["mark_bg"], C["mark_text"]
BLANK_RE = re.compile(r"(_{2,}(?::_{2,})?(?:/_{2,})?)")

RPR_ORDER = ["ln", "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill",
             "effectLst", "effectDag", "highlight", "uLnTx", "uLn", "uFillTx", "uFill",
             "latin", "ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst"]

log = []


def slide(n):
    return prs.slides[n - 1]


def shape_by_id(sl, sid):
    for s in sl.shapes:
        if s.shape_id == sid:
            return s
    raise KeyError(sid)


def sort_rpr(rpr):
    kids = list(rpr)
    kids.sort(key=lambda e: RPR_ORDER.index(etree.QName(e).localname)
              if etree.QName(e).localname in RPR_ORDER else 99)
    for k in kids:
        rpr.remove(k)
    for k in kids:
        rpr.append(k)


def get_rpr(r):
    rpr = r.find(A + "rPr")
    if rpr is None:
        rpr = etree.Element(A + "rPr")
        r.insert(0, rpr)
    return rpr


def set_color(r, hexval):
    rpr = get_rpr(r)
    for tag in ("noFill", "solidFill", "gradFill"):
        for f in rpr.findall(A + tag):
            rpr.remove(f)
    sf = etree.SubElement(rpr, A + "solidFill")
    etree.SubElement(sf, A + "srgbClr").set("val", hexval)
    sort_rpr(rpr)


def mark_run(r):
    set_color(r, MARK_FG)
    rpr = get_rpr(r)
    for h in rpr.findall(A + "highlight"):
        rpr.remove(h)
    hl = etree.SubElement(rpr, A + "highlight")
    etree.SubElement(hl, A + "srgbClr").set("val", MARK_BG)
    sort_rpr(rpr)


def clone_run(tpl_r, text):
    r = copy.deepcopy(tpl_r)
    for h in r.findall(A + "rPr/" + A + "hlinkClick"):
        h.getparent().remove(h)
    r.find(A + "t").text = text
    return r


def rebuild_paragraph(p, tpl_runs, segs):
    """segs: list of [style, text(, url)] where style is r<i> (keep run i format) or m<i> (marker on run i)."""
    pel = p._p
    for r in pel.findall(A + "r"):
        pel.remove(r)
    end = pel.find(A + "endParaRPr")
    for seg in segs:
        style, text = seg[0], seg[1]
        idx = int(style[1:])
        r = clone_run(tpl_runs[idx], text)
        if style[0] == "m":
            mark_run(r)
        opt = seg[2] if len(seg) > 2 else None
        if isinstance(opt, str):
            opt = {"url": opt}
        opt = opt or {}
        if "size" in opt:
            get_rpr(r).set("sz", str(int(round(opt["size"] * 100))))
        if "bold" in opt:
            get_rpr(r).set("b", "1" if opt["bold"] else "0")
        if "color" in opt:
            set_color(r, opt["color"])
        if end is not None:
            end.addprevious(r)
        else:
            pel.append(r)
        if opt.get("url"):
            _Run(r, p).hyperlink.address = opt["url"]


def auto_blank_segs(tpl_runs):
    segs = []
    for i, r in enumerate(tpl_runs):
        t = r.find(A + "t").text or ""
        for part in BLANK_RE.split(t):
            if not part:
                continue
            if BLANK_RE.fullmatch(part):
                segs.append(["m%d" % i, MARK])
            else:
                if segs and segs[-1][0][0] == "m" and not part[0].isspace():
                    part = " " + part  # keep a gap after the marker
                segs.append(["r%d" % i, part])
    return segs


def para_runs(p):
    return [copy.deepcopy(r) for r in p._p.findall(A + "r")]


def textwidth_in(text, pt):
    w = 0.0
    for ch in text:
        if "가" <= ch <= "힣":
            w += 1.0
        elif ch == " ":
            w += 0.28
        else:
            w += 0.58
    return w * pt / 72.0


def add_textbox(sl, x, y, w, h, runs, size, align=None, anchor=MSO_ANCHOR.TOP, fill=None,
                wrap=True, inset=0):
    tb = sl.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Emu(inset))
    if fill:
        tb.fill.solid()
        tb.fill.fore_color.rgb = RGBColor.from_string(fill)
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if align is not None:
            p.alignment = align
        for text, font, color, psize, mark in para:
            r = p.add_run()
            r.text = text
            r.font.name = font
            r.font.size = Pt(psize or size)
            r.font.color.rgb = RGBColor.from_string(color)
            rpr = r._r.get_or_add_rPr()
            # East Asian font too, so Hangul uses the embedded Pretendard
            ea = etree.SubElement(rpr, A + "ea")
            ea.set("typeface", font)
            sort_rpr(rpr)
            if mark:
                mark_run(r._r)
    return tb


def rr_geom(shape, radius, w, h):
    """Turn a text box's rect geometry into a rounded rectangle (E6-pres5, UI token radius 8)."""
    sp_pr = shape._element.spPr
    for g in sp_pr.findall(A + "prstGeom"):
        g.set("prst", "roundRect")
        av = g.find(A + "avLst")
        if av is None:
            av = etree.SubElement(g, A + "avLst")
        for gd in list(av):
            av.remove(gd)
        gd = etree.SubElement(av, A + "gd")
        gd.set("name", "adj")
        gd.set("fmla", "val %d" % int(50000 * radius / (min(w, h) / 2.0)))


prs = Presentation(ORIG)

# 1. explicit run rewrites
def set_paragraphs(tf, tpl_runs, paras):
    """Rewrite a text frame as len(paras) paragraphs, each cloned from the first one (pPr kept)."""
    ps = tf.paragraphs
    first = ps[0]._p
    for extra in ps[1:]:
        extra._p.getparent().remove(extra._p)
    anchor = first
    for k, segs in enumerate(paras):
        if k == 0:
            pel = first
        else:
            pel = copy.deepcopy(first)
            anchor.addnext(pel)
            anchor = pel
        from pptx.text.text import _Paragraph
        rebuild_paragraph(_Paragraph(pel, tf), tpl_runs, segs)


VARIANT = C.get("variant", "now")  # one-line switch in content.json, e.g. "after_E4-L1f"
for spec in C["runs"]:
    sl = slide(spec["slide"])
    if VARIANT in spec.get("variants", {}):
        spec = dict(spec, segs=spec["variants"][VARIANT])
    if spec.get("table"):
        tbl = next(s for s in sl.shapes if s.has_table).table
        tf = tbl.cell(spec["row"], spec["col"]).text_frame
    elif spec.get("tshape"):
        tbl = shape_by_id(sl, spec["tshape"]).table
        tf = tbl.cell(spec["row"], spec["col"]).text_frame
    else:
        sh = shape_by_id(sl, spec["shape"])
        tf = sh.text_frame
        for k in ("x", "y", "w", "h"):
            if k in spec.get("geom", {}):
                setattr(sh, {"x": "left", "y": "top", "w": "width", "h": "height"}[k], Emu(spec["geom"][k]))
    before = tf.text
    pi = spec.get("para", 0)  # rewrite only this paragraph (default: the first)
    tpl = para_runs(tf.paragraphs[pi])
    if not tpl and spec.get("tpl_from"):
        r_, c_ = spec["tpl_from"]
        tpl = para_runs(tbl.cell(r_, c_).text_frame.paragraphs[0])
    if "paras" in spec:
        set_paragraphs(tf, tpl, spec["paras"])
    else:
        rebuild_paragraph(tf.paragraphs[pi], tpl, spec["segs"])
    log.append(("runs", spec["slide"], spec.get("shape", spec.get("tshape", "table")),
                before, tf.text))

# 2. automatic blank -> marker
for spec in C["auto_blanks"]:
    sh = shape_by_id(slide(spec["slide"]), spec["shape"])
    for p in sh.text_frame.paragraphs:
        tpl = para_runs(p)
        if not any(BLANK_RE.search(r.find(A + "t").text or "") for r in tpl):
            continue
        before = p.text
        rebuild_paragraph(p, tpl, auto_blank_segs(tpl))
        log.append(("auto", spec["slide"], spec["shape"], before, p.text))

# 3. table cells -> marker (+ unit in original style)
for spec in C["table_cells"]:
    sh = shape_by_id(slide(spec["slide"]), spec["shape"])
    tbl = sh.table
    tr, tc = spec["tpl_cell"]
    tpl_run = para_runs(tbl.cell(tr, tc).text_frame.paragraphs[0])[0]
    for row, col, unit in spec["cells"]:
        p = tbl.cell(row, col).text_frame.paragraphs[0]
        runs = para_runs(p) or [copy.deepcopy(tpl_run)]
        before = p.text
        segs = [["m0", MARK]] + ([["r0", unit]] if unit else [])
        rebuild_paragraph(p, runs, segs)
        log.append(("cell", spec["slide"], (row, col), before, p.text))

# 4. remove placeholder texts inside screenshot frames
for spec in C["remove_shapes"]:
    sl = slide(spec["slide"])
    for sid in spec["shapes"]:
        sh = shape_by_id(sl, sid)
        log.append(("remove", spec["slide"], sid, sh.text_frame.text, ""))
        sh._element.getparent().remove(sh._element)

# 5. screenshots + honesty stamps
os.makedirs(IMGDIR, exist_ok=True)
for i, spec in enumerate(C["pictures"], 1):
    sl = slide(spec["slide"])
    if spec.get("placeholder"):  # E6-pres5: capture slot to be replaced later (dashed box + note)
        ph = add_textbox(sl, spec["x"], spec["y"], spec["w"], spec["h"],
                         [[(spec["placeholder"], "Pretendard SemiBold", "52525B", 11, False)]], 11,
                         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, fill="F7F7F5", inset=182880)
        rr_geom(ph, 50800, spec["w"], spec["h"])  # UI 토큰: --bg 바탕, --ink-3 점선, 모서리 8
        ph.line.color.rgb = RGBColor.from_string("8B8B94")
        ph.line.width = Pt(0.75)
        ph.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        log.append(("placeholder", spec["slide"], "", "", spec["placeholder"]))
        continue
    im = Image.open(spec["src"]).convert("RGB").crop(tuple(spec["crop"]))
    out = os.path.join(IMGDIR, "s11_shot%d.png" % i)
    im.save(out)
    pic = sl.shapes.add_picture(out, Emu(spec["x"]), Emu(spec["y"]), Emu(spec["w"]), Emu(spec["h"]))
    pic.line.color.rgb = RGBColor.from_string("E4E4E7")  # UI 토큰 --line
    pic.line.width = Pt(0.75)
    log.append(("picture", spec["slide"], os.path.basename(spec["src"]), spec["crop"], spec.get("stamp", "")))
    if not spec.get("stamp"):
        continue
    size = 8
    w = int((textwidth_in(spec["stamp"], size) + 0.22) * 914400)
    h = 228600
    add_textbox(sl, spec["x"] + 54864, spec["y"] + spec["h"] - h - 54864, w, h,
                [[(spec["stamp"], "Pretendard SemiBold", "FFFFFF", size, False)]], size,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, fill="B5171F", wrap=False, inset=45720)

# 6. cover: prototype / demo video placeholders with QR boxes
QR = {}
_qr_path = os.path.join(os.path.dirname(os.path.abspath(CONTENT)), C.get("qr_file", "qr.json"))
if os.path.exists(_qr_path):
    QR = json.load(open(_qr_path, encoding="utf-8"))
cv = C["cover"]
sl = slide(cv["slide"])
for it in cv["items"]:
    q_ = QR.get(it.get("key", ""))
    if q_:
        pic = sl.shapes.add_picture(q_["png"], Emu(it["qr_x"]), Emu(cv["y"]), Emu(cv["qr"]), Emu(cv["qr"]))
        short = q_["url"].split("://", 1)[-1].rstrip("/")
        add_textbox(sl, it["label_x"], cv["y"], it["label_w"], cv["qr"],
                    [[(it["title"], "Pretendard SemiBold", "FFFFFF", 12, False)],
                     [(short, "Pretendard", "D5DAE0", 8, False)]],
                    12, anchor=MSO_ANCHOR.MIDDLE)
        log.append(("cover-qr", 1, it["title"], "", q_["url"]))
        continue
    q = sl.shapes.add_textbox(Emu(it["qr_x"]), Emu(cv["y"]), Emu(cv["qr"]), Emu(cv["qr"]))
    q.line.color.rgb = RGBColor.from_string("AAB4C0")
    q.line.width = Pt(0.75)
    q.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    tf = q.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.auto_size = MSO_AUTO_SIZE.NONE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = cv["qr_text"]
    r.font.name = "Pretendard SemiBold"
    etree.SubElement(r._r.get_or_add_rPr(), A + "ea").set("typeface", "Pretendard SemiBold")
    sort_rpr(r._r.get_or_add_rPr())
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string("AAB4C0")
    add_textbox(sl, it["label_x"], cv["y"], it["label_w"], cv["qr"],
                [[(it["title"], "Pretendard SemiBold", "FFFFFF", 12, False)],
                 [(it["sub"] + "  ", "Pretendard", "AAB4C0", 9, False),
                  (cv["value_mark"], "Pretendard SemiBold", MARK_FG, 9, True)]],
                12, anchor=MSO_ANCHOR.MIDDLE)
    log.append(("cover", 1, it["title"], "", cv["value_mark"]))

# 7. slide 11: prototype / video line in the scenario box header
ll = C["link_line"]
if QR and ll.get("segs_after_qr"):
    ll = dict(ll, segs=ll["segs_after_qr"])
colors = {"gray": ("Pretendard SemiBold", "6A6F75", False), "dark": ("Pretendard SemiBold", "14161A", False),
          "mark": ("Pretendard SemiBold", MARK_FG, True), "muted": ("Pretendard", "6A6F75", False)}
runs = [(t, colors[k][0], colors[k][1], 10, colors[k][2]) for k, t in ll["segs"]]
add_textbox(slide(ll["slide"]), ll["x"], ll["y"], ll["w"], ll["h"], [runs], 10, align=PP_ALIGN.RIGHT)
log.append(("linkline", ll["slide"], "", "", "".join(t for _, t in ll["segs"])))

# 7b. extra text boxes (provenance lines etc.)
# "segs": one paragraph; or "paras": several. A seg is [kind, text] or [kind, text, {size, color, font}].
for tb in C.get("textboxes", []):
    paras = tb.get("paras") or [tb["segs"]]
    runs_all = []
    for segs in paras:
        pr = []
        for seg in segs:
            k, t = seg[0], seg[1]
            opt = seg[2] if len(seg) > 2 else {}
            font = opt.get("font") or (colors[k][0] if k in colors else "Pretendard")
            col = opt.get("color") or (colors[k][1] if k in colors else k)
            pr.append((t, font, col, opt.get("size", tb.get("size", 9)), k == "mark"))
        runs_all.append(pr)
    anchor = MSO_ANCHOR.TOP if tb.get("anchor") == "top" else MSO_ANCHOR.MIDDLE
    align = {"right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}.get(tb.get("align"))
    add_textbox(slide(tb["slide"]), tb["x"], tb["y"], tb["w"], tb["h"], runs_all, tb.get("size", 9),
                align=align, anchor=anchor)
    log.append(("textbox", tb["slide"], "", "", " / ".join("".join(s[1] for s in segs) for segs in paras)))

for g in C.get("geoms", []):
    sh = find_shape(slide(g["slide"]), g["shape"]) if "find_shape" in globals() else shape_by_id(slide(g["slide"]), g["shape"])
    for k in ("x", "y", "w", "h"):
        if k in g:
            setattr(sh, {"x": "left", "y": "top", "w": "width", "h": "height"}[k], Emu(g[k]))
    log.append(("geom", g["slide"], g["shape"], "", str(g)))

# 7c. "설계" page badges and inline "예정" tags (design pages: what the screen does not show yet)
def find_shape(sl, sid):
    def walk(shapes):
        for s_ in shapes:
            if s_.shape_id == sid:
                return s_
            if s_.shape_type == 6:
                got = walk(s_.shapes)
                if got is not None:
                    return got
        return None
    got = walk(sl.shapes)
    if got is None:
        raise KeyError(sid)
    return got


TAG = C.get("tag_style", {"color": "4F555C", "bg": "E2E5E0", "scale": 0.72, "min": 7})


def run_size(r, default=10.0):
    rpr = r.find(A + "rPr")
    if rpr is not None and rpr.get("sz"):
        return int(rpr.get("sz")) / 100.0
    return default


for tg in C.get("tags", []):
    sl = slide(tg["slide"])
    if "tshape" in tg:
        tf = find_shape(sl, tg["tshape"]).table.cell(tg["row"], tg["col"]).text_frame
    else:
        tf = find_shape(sl, tg["shape"]).text_frame
    done = False
    for p_ in tf.paragraphs:
        for r in p_._p.findall(A + "r"):
            t = r.find(A + "t").text or ""
            k = t.find(tg["after"])
            if k < 0:
                continue
            end = k + len(tg["after"])
            head, tail = t[:end], t[end:]
            r.find(A + "t").text = head
            size = run_size(r, tg.get("base", 10.0))
            gap = clone_run(r, " ")
            tag = clone_run(r, " %s " % tg.get("text", "예정"))
            trpr = get_rpr(tag)
            trpr.set("sz", str(int(round(max(TAG["min"], size * TAG["scale"]) * 100))))
            trpr.set("b", "0")
            for u in trpr.findall(A + "uLn") + trpr.findall(A + "uFill"):
                trpr.remove(u)
            trpr.set("u", "none")
            set_color(tag, TAG["color"])
            for h in trpr.findall(A + "highlight"):
                trpr.remove(h)
            hl = etree.SubElement(trpr, A + "highlight")
            etree.SubElement(hl, A + "srgbClr").set("val", TAG["bg"])
            sort_rpr(trpr)
            r.addnext(gap)
            gap.addnext(tag)
            if tail:
                tag.addnext(clone_run(r, tail))
            done = True
            break
        if done:
            break
    if not done:
        raise ValueError("tag target not found: %r" % (tg,))
    log.append(("tag", tg["slide"], tg.get("shape", tg.get("tshape")), tg["after"], tg.get("text", "예정")))

for bd in C.get("badges", []):
    sl = slide(bd["slide"])
    tb = add_textbox(sl, bd["x"], bd["y"], bd["w"], bd["h"],
                     [[(bd["text"], "Pretendard SemiBold", bd.get("color", "6A6F75"), bd.get("size", 8.5), False)]],
                     bd.get("size", 8.5), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    tb.line.color.rgb = RGBColor.from_string(bd.get("line", "9AA0A6"))
    tb.line.width = Pt(0.75)
    log.append(("badge", bd["slide"], "", "", bd["text"]))

# 7c2. bordered flow boxes (E6-pres5: full-flow strips on slides 6·9). Same seg format as "textboxes".
#      {"slide", "x","y","w","h", "paras": [[seg...]], "size", "line": hex|None, "line_w": pt,
#       "dash": bool, "fill": hex|None, "anchor": "top"|"middle", "align": "center"|"right"|None, "inset": emu}
for bx in C.get("boxes", []):
    runs_all = []
    for segs in bx["paras"]:
        pr = []
        for seg in segs:
            k, t = seg[0], seg[1]
            opt = seg[2] if len(seg) > 2 else {}
            font = opt.get("font") or (colors[k][0] if k in colors else "Pretendard")
            col = opt.get("color") or (colors[k][1] if k in colors else k)
            pr.append((t, font, col, opt.get("size", bx.get("size", 9)), k == "mark"))
        runs_all.append(pr)
    anchor = MSO_ANCHOR.TOP if bx.get("anchor") == "top" else MSO_ANCHOR.MIDDLE
    align = {"right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}.get(bx.get("align"))
    tb = add_textbox(slide(bx["slide"]), bx["x"], bx["y"], bx["w"], bx["h"], runs_all, bx.get("size", 9),
                     align=align, anchor=anchor, fill=bx.get("fill"), inset=bx.get("inset", 45720))
    if bx.get("radius"):  # UI 토큰 모서리 8: 사각 글상자를 둥근 사각형으로 바꾼다
        rr_geom(tb, bx["radius"], bx["w"], bx["h"])
    if bx.get("line"):
        tb.line.color.rgb = RGBColor.from_string(bx["line"])
        tb.line.width = Pt(bx.get("line_w", 0.75))
        if bx.get("dash"):
            tb.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        tb.line.fill.background()
    log.append(("box", bx["slide"], "", "", " / ".join("".join(s[1] for s in segs) for segs in bx["paras"])))

# 7d. table column widths / row heights (E6-pres3: wider "현장 측정" column on slide 12)
for tg in C.get("table_geoms", []):
    tbl = shape_by_id(slide(tg["slide"]), tg["shape"]).table
    for i, w in enumerate(tg.get("cols", [])):
        if w:
            tbl.columns[i].width = Emu(w)
    for i, h in tg.get("rows", {}).items():
        tbl.rows[int(i)].height = Emu(h)
    log.append(("table_geom", tg["slide"], tg["shape"], "", json.dumps(tg)))

# 8. speaker notes
for spec in C["notes"]:
    sl = slide(spec["slide"])
    tf = sl.notes_slide.notes_text_frame
    txt = tf.text
    before = txt
    if "set" in spec:
        txt = spec["set"]
    if spec.get("blanks"):
        txt = re.sub(r"_{2,}", C["note_mark"], txt)
        txt = txt.replace(C["note_tag_old"], C["note_tag_new"])
    for a, b in spec.get("replace", []):
        txt = txt.replace(a, b)
    if spec.get("append"):
        txt = txt.rstrip() + " " + spec["append"]
    # keep the original run properties (lang only in the kit) on every rewritten paragraph
    first_r = tf._txBody.find(".//" + A + "r")
    tpl_rpr = copy.deepcopy(first_r.find(A + "rPr")) if first_r is not None and first_r.find(A + "rPr") is not None else None
    tf.text = txt
    if tpl_rpr is not None:
        for r in tf._txBody.iter(A + "r"):
            old = r.find(A + "rPr")
            if old is not None:
                r.remove(old)
            r.insert(0, copy.deepcopy(tpl_rpr))
    log.append(("notes", spec["slide"], "", before, txt))

# 9. drop whole slides last, so every slide number above stays the kit's number (E6-pres5).
#    Page numbers on the slides are fields and renumber themselves; the contents page is text (slide 2 runs).
_ids = prs.slides._sldIdLst
_all = list(_ids)
for n in sorted(C.get("delete_slides", []), reverse=True):
    el = _all[n - 1]
    title = next((s.text_frame.text for s in prs.slides[n - 1].shapes
                  if s.has_text_frame and s.shape_id == 4), "")
    prs.part.drop_rel(el.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
    _ids.remove(el)
    log.append(("delete_slide", n, "", title, ""))

prs.save(OUT)
with open(os.path.join(os.path.dirname(OUT), "build_log.json"), "w", encoding="utf-8") as f:
    json.dump(log, f, ensure_ascii=False, indent=1, default=str)
print("saved", OUT, "edits", len(log))
